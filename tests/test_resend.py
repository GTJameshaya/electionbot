"""Тесты «Отправить в чат» из админки и связи с CRM (без сети: Telegram и CRM подменены).

Запуск:  python -m unittest discover -s tests -v
"""
import asyncio
import io
import json
import os
import sys
import unittest
from datetime import timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
os.environ["HOUSES_FILE"] = str(ROOT / "tests" / "fixtures" / "houses.json")

import httpx  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import inspect, text  # noqa: E402

import admin  # noqa: E402
import bot  # noqa: E402
import config  # noqa: E402
import crm_ingest  # noqa: E402
from captions import build_caption  # noqa: E402
from config import AdminSettings  # noqa: E402
from crm_ingest import CrmConfig, TicketRequest  # noqa: E402
from db import ChatResend, Database, Submission, current_period, utcnow  # noqa: E402
from test_admin import ADMIN, TOKEN, USER, csrf_of  # noqa: E402
from test_offline import Base  # noqa: E402

CHAT = -1009990001
TAG = "@test_crm_bot"
CRM = CrmConfig(url="https://crm.example.test/api/ingest/ticket", token="fake-token-for-tests")
OLD = "2024-05"   # архивный месяц с показаниями


# ------------------------------------------------------------------ подделки
class FakeSender:
    """Вместо Telegram: запоминает отправки, возвращает message_id (или падает)."""

    def __init__(self, fail=None, log=None):
        self.calls, self.fail, self.log = [], fail, log if log is not None else []
        self.next_id = 500

    async def send_document(self, chat_id, data, filename, caption):
        self.log.append("telegram")
        if self.fail:
            raise self.fail
        self.next_id += 1
        self.calls.append(dict(chat_id=chat_id, data=data, filename=filename, caption=caption,
                               message_id=self.next_id))
        return self.next_id


def parse_multipart(request: httpx.Request) -> dict:
    """{поле: [значение | (имя файла, байты)]} из multipart-запроса."""
    ctype = request.headers["content-type"]
    boundary = ctype.split("boundary=")[1].encode()
    out = {}
    for part in request.content.split(b"--" + boundary):
        if b"\r\n\r\n" not in part:
            continue
        head, body = part.split(b"\r\n\r\n", 1)
        body = body[:-2] if body.endswith(b"\r\n") else body
        head = head.decode()
        name = head.split('name="')[1].split('"')[0]
        if 'filename="' in head:
            value = (head.split('filename="')[1].split('"')[0], body)
        else:
            value = body.decode()
        out.setdefault(name, []).append(value)
    return out


class FakeCrm:
    """Вместо CRM: httpx.MockTransport. responses — список (код, json) или исключений по порядку."""

    def __init__(self, *responses, log=None):
        self.responses = list(responses) or [(200, {"ticket_number": 12, "ticket_id": 40, "created": True})]
        self.requests, self.log = [], log if log is not None else []

    def handler(self, request):
        self.log.append("crm")
        self.requests.append(dict(headers=request.headers, url=str(request.url), form=parse_multipart(request)))
        r = self.responses.pop(0) if len(self.responses) > 1 else self.responses[0]
        if isinstance(r, Exception):
            raise r
        return httpx.Response(r[0], json=r[1])

    def patch(self):
        return patch("crm_ingest.make_client",
                     lambda cfg: httpx.AsyncClient(transport=httpx.MockTransport(self.handler), timeout=cfg.timeout))

    def field(self, i, name):
        return self.requests[i]["form"].get(name, [None])[0]


def no_crm():
    def boom(cfg):
        raise AssertionError("CRM не должна вызываться")
    return patch("crm_ingest.make_client", boom)


# ------------------------------------------------------------------ админка
class ResendBase(Base):
    crm = CRM
    chat = CHAT

    def setUp(self):
        super().setUp()
        self.db.seed_roles([ADMIN], [USER])
        self.db.touch_user(ADMIN, "boss", "Анна", None)
        self.db.save_reading(1, "1", current_period(), 13100, None)
        self.db.save_reading(1, "1", OLD, 12000, None)
        sub = self.db.subscribers(1)[0][0]
        self.db.save_reading(1, sub, OLD, 500, None, kind="commerce")
        self.sender = FakeSender()
        self.settings = AdminSettings(secret_key="s" * 40, bot_token=TOKEN, group_chat_id=self.chat,
                                      target_bot_tag=TAG, crm=self.crm)
        self.client = TestClient(admin.create_app(self.db, self.settings, sender=self.sender))
        self.login()

    def login(self, uid=ADMIN):
        raw = self.db.create_login_token(uid)
        self.client.get(f"/auth/tg?token={raw}", follow_redirects=False)

    def csrf(self):
        return csrf_of(self.client.get("/").text)

    def confirm(self, house=1, kind="residential", period=OLD, mode="plain", ticket=None, csrf=None):
        data = {"csrf": self.csrf() if csrf is None else csrf, "house": house, "kind": kind,
                "period": period, "mode": mode}
        if ticket:
            data["ticket"] = ticket
        return self.client.post("/resend/confirm", data=data, follow_redirects=False)

    def send(self, token, csrf=None):
        return self.client.post("/resend/send", data={"csrf": self.csrf() if csrf is None else csrf,
                                                      "token": token}, follow_redirects=False)

    def confirm_and_send(self, **kw):
        r = self.confirm(**kw)
        self.assertEqual(r.status_code, 200, r.text[:300])
        token = token_of(r.text)
        r = self.send(token)
        self.assertEqual(r.status_code, 303)
        return token, self.client.get("/resend").text

    def rows(self):
        with self.db.Session() as s:
            return list(s.query(ChatResend).order_by(ChatResend.id))


def token_of(html):
    import re
    return re.search(r'name="token" value="([^"]+)"', html).group(1)


class TestResendPage(ResendBase):
    def test_page_lists_houses_kinds_and_periods(self):
        html = self.client.get("/resend").text
        self.assertIn("Отправить в чат", html)
        for name in ("Дом 1", "Дом 2", "Жилые", "Коммерция", "Май 2024"):
            self.assertIn(name, html)
        self.assertNotIn("CRM не настроена", html)
        # у дома 1: текущий месяц + архивный с показаниями; у дома 2 — только текущий
        self.assertEqual(html.count('name="period" value="2024-05"'), 2)
        self.assertIn('href="/resend"', self.client.get("/").text)  # пункт меню

    def test_confirm_shows_details(self):
        r = self.confirm(kind="commerce", mode="ticket")
        self.assertEqual(r.status_code, 200)
        for s in ("Дом 1", "Май 2024", "Коммерция", "С тикетом", "новый тикет", str(CHAT), TAG):
            self.assertIn(s, r.text)
        self.assertEqual(self.sender.calls, [])
        self.assertEqual(self.rows()[0].status, "pending")
        self.assertNotIn("Анна", self.client.get("/resend").text.split("Журнал")[1])  # pending не в журнале

    def test_period_without_readings_rejected(self):
        r = self.confirm(period="2023-01")
        self.assertEqual(r.status_code, 303)
        self.assertEqual(self.rows(), [])
        r = self.confirm(house=99)
        self.assertEqual(self.rows(), [])
        r = self.confirm(mode="hack")
        self.assertEqual(self.rows(), [])


class TestResendPlain(ResendBase):
    def test_residential_plain_no_tag_no_crm(self):
        with no_crm():
            _, page = self.confirm_and_send(kind="residential", period=OLD, mode="plain")
        self.assertEqual(len(self.sender.calls), 1)
        call = self.sender.calls[0]
        self.assertEqual(call["chat_id"], CHAT)
        self.assertEqual(call["caption"], build_caption("", 1, OLD))
        self.assertTrue(call["caption"].startswith("Привет! Показания за текущий 05.2024 — Дом 1"))
        self.assertNotIn(TAG, call["caption"])
        self.assertEqual(call["filename"], "Дом_1_Май_2024.xlsx")     # месяц отчёта, не дата отправки
        ws = load_workbook(io.BytesIO(call["data"])).active
        self.assertIn(12000, [c.value for row in ws.iter_rows() for c in row])
        row = self.rows()[0]
        self.assertEqual((row.status, row.mode, row.message_id, row.requested_by, row.crm_status),
                         ("sent", "plain", 501, ADMIN, None))
        self.assertIn("Отправлено в чат", page)
        self.assertIn("не нужно", page)
        with self.db.Session() as s:
            sub = s.query(Submission).one()
        self.assertEqual((sub.source, sub.kind, sub.period, sub.submitted_by), ("admin", "residential", OLD, ADMIN))

    def test_commerce_plain(self):
        with no_crm():
            self.confirm_and_send(kind="commerce", period=OLD, mode="plain")
        call = self.sender.calls[0]
        self.assertEqual(call["caption"], build_caption("", 1, OLD, "commerce"))
        self.assertIn("Коммерция", call["caption"])
        self.assertEqual(call["filename"], "Дом_1_Коммерция_Май_2024.xlsx")
        self.assertEqual(self.rows()[0].kind, "commerce")

    def test_send_failure_logged_no_crm(self):
        self.sender.fail = RuntimeError("Telegram недоступен")
        with no_crm():
            _, page = self.confirm_and_send(mode="ticket")
        row = self.rows()[0]
        self.assertEqual(row.status, "failed")
        self.assertIn("Telegram недоступен", row.error)
        self.assertIn("Telegram недоступен", page)
        self.assertIsNone(self.db.get_crm_ticket(1, OLD, "residential"))


class TestResendTicket(ResendBase):
    def test_ticket_mode_tag_and_crm_metadata(self):
        crm = FakeCrm()
        with crm.patch():
            _, page = self.confirm_and_send(kind="commerce", period=OLD, mode="ticket")
        call = self.sender.calls[0]
        self.assertEqual(call["caption"], build_caption(TAG, 1, OLD, "commerce"))
        self.assertTrue(call["caption"].startswith(TAG + " Привет!"))
        self.assertEqual(len(crm.requests), 1)
        req = crm.requests[0]
        self.assertEqual(req["url"], CRM.url)
        self.assertEqual(req["headers"]["authorization"], "Bearer fake-token-for-tests")
        self.assertEqual(crm.field(0, "text"), call["caption"])
        self.assertEqual(crm.field(0, "chat_id"), str(CHAT))
        self.assertEqual(crm.field(0, "source_message_id"), str(call["message_id"]))
        self.assertEqual(crm.field(0, "author"), "Анна (@boss)")
        self.assertIsNone(crm.field(0, "reply_to_ticket"))
        self.assertIsNone(crm.field(0, "category"))
        meta = json.loads(crm.field(0, "meta"))
        self.assertEqual(meta, {"house": "Дом 1", "house_id": 1, "period": OLD, "kind": "commerce",
                                "source": "admin_resend", "user_id": ADMIN})
        fname, data = crm.field(0, "files")
        self.assertEqual((fname, data), (call["filename"], call["data"]))
        row = self.rows()[0]
        self.assertEqual((row.crm_status, row.crm_ticket_number, row.crm_created), ("ok", 12, True))
        self.assertIn("тикет #12 создан", page)
        t = self.db.get_crm_ticket(1, OLD, "commerce")
        self.assertEqual((t.ticket_number, t.ticket_id, t.source), (12, 40, "admin_resend"))

    def test_append_is_default_when_ticket_exists(self):
        self.db.save_crm_ticket(1, OLD, "residential", 7, 30, "bot")
        html = self.client.get("/resend").text
        self.assertIn('<option value="append" selected>Дописать в тикет #7</option>', html)
        r = self.confirm(mode="ticket", ticket="append")
        self.assertIn("дописать в тикет #7", r.text)
        crm = FakeCrm((200, {"ticket_number": 7, "ticket_id": 30, "created": False}))
        with crm.patch():
            self.send(token_of(r.text))
        self.assertEqual(crm.field(0, "reply_to_ticket"), "7")
        self.assertEqual(self.rows()[0].reply_to_ticket, 7)
        self.assertIn("тикет #7 дополнен", self.client.get("/resend").text)
        # «Новый тикет» — без reply_to_ticket, сохранённый номер обновляется
        crm = FakeCrm((200, {"ticket_number": 15, "ticket_id": 41, "created": True}))
        with crm.patch():
            self.confirm_and_send(mode="ticket", ticket="new")
        self.assertIsNone(crm.field(0, "reply_to_ticket"))
        self.assertEqual(self.db.get_crm_ticket(1, OLD, "residential").ticket_number, 15)

    def test_reply_404_falls_back_to_new_ticket(self):
        self.db.save_crm_ticket(1, OLD, "residential", 7, 30, "bot")
        crm = FakeCrm((404, {"detail": "Тикет #7 не найден"}),
                      (200, {"ticket_number": 21, "ticket_id": 50, "created": True}))
        with crm.patch():
            _, page = self.confirm_and_send(mode="ticket", ticket="append")
        self.assertEqual(len(crm.requests), 2)
        self.assertEqual(crm.field(0, "reply_to_ticket"), "7")
        self.assertIsNone(crm.field(1, "reply_to_ticket"))
        self.assertEqual(self.db.get_crm_ticket(1, OLD, "residential").ticket_number, 21)
        self.assertIn("старый тикет не найден — создан новый", page)
        self.assertEqual(len(self.sender.calls), 1)  # в Telegram — один раз

    def test_warnings_shown(self):
        crm = FakeCrm((200, {"ticket_number": 3, "ticket_id": 3, "created": True,
                             "warnings": ["не удалось написать в чат"]}))
        with crm.patch(), self.assertLogs("crm", "WARNING") as logs:
            r = self.send(token_of(self.confirm(mode="ticket").text))
            flash_page = self.client.get(r.headers["location"]).text
        self.assertIn('class="flash warn"', flash_page)
        self.assertIn("предупреждения: не удалось написать в чат", flash_page)
        self.assertIn("не удалось написать в чат", self.rows()[0].crm_message)
        self.assertTrue(any("не удалось написать в чат" in m for m in logs.output))

    def test_crm_errors_do_not_break_send(self):
        for resp, expected in [((409, {"detail": "chat detached"}), "чат отключён от CRM"),
                               (httpx.ConnectError("refused"), "не удалось связаться с CRM")]:
            crm = FakeCrm(resp)
            with crm.patch():
                _, page = self.confirm_and_send(mode="ticket")
            row = self.rows()[-1]
            self.assertEqual((row.status, row.crm_status), ("sent", "error"))
            self.assertIn(expected, row.crm_message)
            self.assertIn(expected, page)
        self.assertEqual(len(self.sender.calls), 2)
        self.assertIsNone(self.db.get_crm_ticket(1, OLD, "residential"))


class TestResendSafety(ResendBase):
    def test_double_submit_sends_once(self):
        r = self.confirm()
        token = token_of(r.text)
        self.assertEqual(self.send(token).status_code, 303)
        r = self.send(token)
        self.assertEqual(r.status_code, 303)
        self.assertIn("уже выполнена", self.client.get("/resend").text)
        self.assertEqual(len(self.sender.calls), 1)
        self.assertEqual(len(self.rows()), 1)
        self.assertIn("data-once", self.confirm().text)  # и кнопка гаснет в браузере

    def test_stale_sending_marked_unknown(self):
        old = self.db.claim_resend(token_of(self.confirm().text), ADMIN)
        fresh_token = token_of(self.confirm().text)
        with patch("db.utcnow", return_value=utcnow() - timedelta(minutes=11)):
            self.db.claim_resend(token_of(self.confirm().text), ADMIN)   # «зависла» 11 минут назад
        self.db.claim_resend(fresh_token, ADMIN)                         # идёт прямо сейчас
        html = self.client.get("/resend").text
        statuses = {r.id: r.status for r in self.rows()}
        self.assertEqual(sorted(statuses.values()), ["sending", "sending", "unknown"])
        self.assertEqual(statuses[old.id], "sending")
        self.assertIn("неизвестно", html)
        self.assertIn("файл мог дойти, проверьте чат", html)
        self.assertEqual(self.db.expire_stale_resends(10), 0)            # повторно ничего не трогаем

    def test_telegram_timeout_is_unknown_without_crm(self):
        from telegram.error import TimedOut
        for exc in (TimedOut(), httpx.ReadTimeout("slow")):
            self.sender.fail = exc
            with no_crm():
                _, page = self.confirm_and_send(mode="ticket")
            row = self.rows()[-1]
            self.assertEqual(row.status, "unknown")
            self.assertIn("файл мог дойти, проверьте чат", row.error)
            self.assertIsNone(row.crm_status)
            self.assertIn("файл мог дойти, проверьте чат", page)
        with self.db.Session() as s:
            self.assertEqual(s.query(Submission).count(), 0)
        self.sender.fail = RuntimeError("Bad Request: chat not found")   # обычная ошибка — failed
        with no_crm():
            self.confirm_and_send(mode="ticket")
        self.assertEqual(self.rows()[-1].status, "failed")

    def test_token_expired_or_foreign(self):
        token = token_of(self.confirm().text)
        with patch("db.utcnow", return_value=utcnow() + timedelta(minutes=16)):
            self.send(token)
        self.assertEqual(self.sender.calls, [])
        self.assertIsNone(self.db.claim_resend(token_of(self.confirm().text), ADMIN + 1))
        self.assertIsNone(self.db.claim_resend("", ADMIN))

    def test_csrf_missing_or_invalid(self):
        for bad in ("", "wrong"):
            r = self.confirm(csrf=bad)
            self.assertEqual(r.status_code, 403)
        self.assertEqual(self.rows(), [])
        token = token_of(self.confirm().text)
        for bad in ("", "wrong"):
            self.assertEqual(self.send(token, csrf=bad).status_code, 403)
        self.assertEqual(self.sender.calls, [])
        r = self.client.post("/resend/send", data={"token": token}, follow_redirects=False)
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.sender.calls, [])

    def test_non_admin_denied(self):
        token = token_of(self.confirm().text)
        csrf = self.csrf()
        anon = TestClient(admin.create_app(self.db, self.settings, sender=self.sender))
        self.assertEqual(anon.get("/resend", follow_redirects=False).headers["location"], "/login")
        r = anon.post("/resend/send", data={"csrf": csrf, "token": token}, follow_redirects=False)
        self.assertEqual(r.headers["location"], "/login")
        self.db.set_role(USER, "admin", "test")
        self.db.set_role(ADMIN, "user", "test")   # роль сняли — сессия больше не действует
        r = self.send(token, csrf=csrf)
        self.assertEqual(r.headers["location"], "/login")
        r = self.client.post("/resend/confirm", data={"csrf": csrf, "house": 1, "kind": "residential",
                                                      "period": OLD, "mode": "plain"}, follow_redirects=False)
        self.assertEqual(r.headers["location"], "/login")
        self.assertEqual(self.sender.calls, [])
        self.assertEqual(len(self.rows()), 1)


class TestResendWithoutCrm(ResendBase):
    crm = CrmConfig()

    def test_ticket_disabled_plain_works(self):
        html = self.client.get("/resend").text
        self.assertIn("CRM_INGEST_URL и CRM_INGEST_TOKEN", html)
        self.assertIn('<option value="ticket" disabled>С тикетом (CRM не настроена)</option>', html)
        r = self.confirm(mode="ticket")
        self.assertEqual(r.status_code, 303)
        self.assertIn("«С тикетом» выключен", self.client.get("/resend").text)
        self.assertEqual(self.rows(), [])
        with no_crm():
            self.confirm_and_send(mode="plain")
        self.assertEqual(len(self.sender.calls), 1)
        self.assertNotIn(TAG, self.sender.calls[0]["caption"])


class TestResendWithoutChat(ResendBase):
    chat = 0

    def test_disabled(self):
        html = self.client.get("/resend").text
        self.assertIn("GROUP_CHAT_ID не задан", html)
        self.assertIn("<button disabled>Отправить в чат</button>", html)
        self.assertEqual(self.confirm().status_code, 303)
        self.assertEqual(self.rows(), [])


# ------------------------------------------------------------------ бот
class TestBotCrm(Base):
    def ctx(self, crm=CRM, fail=None, log=None):
        log = log if log is not None else []

        async def send_document(*a, **kw):
            log.append("telegram")
            if fail:
                raise fail
            return SimpleNamespace(message_id=777)
        tg = SimpleNamespace(send_document=AsyncMock(side_effect=send_document), username="fixture_bot")
        settings = config.Settings(bot_token="x", group_chat_id=CHAT, target_bot_tag=TAG, crm=crm)
        return SimpleNamespace(application=SimpleNamespace(bot_data={"db": self.db, "settings": settings}),
                               bot=tg), settings

    def send(self, ctx, settings, kind="residential"):
        """_send_ledger и затем ожидание фоновой задачи CRM (в боте она идёт фоном)."""
        async def go():
            res = await bot._send_ledger(ctx, settings, self.db, 1, OLD, kind, 5)
            await bot.wait_background(timeout=5)
            return res
        return asyncio.run(go())

    def test_crm_runs_in_background(self):
        """Ответ человеку не ждёт CRM: _send_ledger возвращается, пока запрос в CRM ещё висит."""
        ctx, settings = self.ctx()
        state = {}

        async def go():
            release = asyncio.Event()
            started = asyncio.Event()

            async def slow(request):
                started.set()
                await release.wait()
                state["form"] = parse_multipart(request)
                return httpx.Response(200, json={"ticket_number": 9, "ticket_id": 9, "created": True})
            with patch("crm_ingest.make_client", lambda cfg: httpx.AsyncClient(
                    transport=httpx.MockTransport(slow), timeout=cfg.timeout)):
                res = await asyncio.wait_for(bot._send_ledger(ctx, settings, self.db, 1, OLD, "residential", 5), 2)
                state["result"] = res
                state["pending"] = len(bot._background)
                state["ticket_before"] = self.db.get_crm_ticket(1, OLD, "residential")
                await asyncio.wait_for(started.wait(), 2)
                release.set()
                await bot.wait_background(timeout=5)
            state["left"] = len(bot._background)
        asyncio.run(go())
        self.assertEqual(state["result"], (True, None))
        self.assertEqual(state["pending"], 1)
        self.assertIsNone(state["ticket_before"])
        self.assertEqual(state["left"], 0)                       # ссылка на задачу убрана
        self.assertEqual(state["form"]["source_message_id"], ["777"])
        self.assertEqual(self.db.get_crm_ticket(1, OLD, "residential").ticket_number, 9)

    def test_background_exception_logged(self):
        ctx, settings = self.ctx()
        with patch("bot.push_ledger_ticket", side_effect=RuntimeError("crm exploded")), \
                self.assertLogs("bot", "ERROR") as logs:
            self.assertEqual(self.send(ctx, settings), (True, None))
        self.assertTrue(any("CRM: не удалось создать тикет" in m and "crm exploded" in m for m in logs.output))
        self.assertTrue(self.db.is_submitted(1, OLD))

        async def boom():
            raise ValueError("background boom")

        async def go():
            bot.spawn_background(boom(), "test-task")
            await bot.wait_background(timeout=5)
            await asyncio.sleep(0)   # дать отработать done-callback
        with self.assertLogs("bot", "ERROR") as logs:
            asyncio.run(go())
        self.assertTrue(any("test-task" in m and "background boom" in m for m in logs.output))
        self.assertEqual(len(bot._background), 0)

    def test_reply_404_disabled_no_retry(self):
        self.db.save_crm_ticket(1, OLD, "residential", 5, 5, "bot")
        ctx, settings = self.ctx()
        crm = FakeCrm((404, {"detail": "Not Found"}))
        with crm.patch():
            self.send(ctx, settings)
        self.assertEqual(len(crm.requests), 1)          # «выключено» — новый тикет не пытаемся создать
        self.assertEqual(self.db.get_crm_ticket(1, OLD, "residential").ticket_number, 5)

    def test_not_configured_no_call(self):
        ctx, settings = self.ctx(crm=CrmConfig(url=CRM.url, token=""))
        with no_crm():
            self.assertEqual(self.send(ctx, settings), (True, None))
        self.assertTrue(self.db.is_submitted(1, OLD))

    def test_called_after_send_new_then_append(self):
        log = []
        ctx, settings = self.ctx(log=log)
        crm = FakeCrm((200, {"ticket_number": 12, "ticket_id": 40, "created": True}),
                      (200, {"ticket_number": 12, "ticket_id": 40, "created": False}), log=log)
        with crm.patch():
            self.assertEqual(self.send(ctx, settings, "commerce"), (True, None))
            self.assertEqual(self.send(ctx, settings, "commerce"), (True, None))
        self.assertEqual(log, ["telegram", "crm", "telegram", "crm"])
        self.assertEqual(crm.field(0, "text"), build_caption(TAG, 1, OLD, "commerce"))
        self.assertEqual(crm.field(0, "source_message_id"), "777")
        self.assertEqual(crm.field(0, "author"), "@fixture_bot")
        self.assertEqual(crm.field(0, "chat_id"), str(CHAT))
        self.assertEqual(json.loads(crm.field(0, "meta")),
                         {"house": "Дом 1", "house_id": 1, "period": OLD, "kind": "commerce",
                          "source": "bot", "user_id": 5})
        self.assertIsNone(crm.field(0, "reply_to_ticket"))
        self.assertEqual(crm.field(1, "reply_to_ticket"), "12")   # второй раз — дописываем
        t = self.db.get_crm_ticket(1, OLD, "commerce")
        self.assertEqual((t.ticket_number, t.source), (12, "bot"))

    def test_send_failure_no_crm(self):
        ctx, settings = self.ctx(fail=RuntimeError("network"))
        with no_crm(), self.assertRaises(RuntimeError):
            self.send(ctx, settings)
        self.assertFalse(self.db.is_submitted(1, OLD))

    def test_crm_failure_does_not_break_send(self):
        ctx, settings = self.ctx()
        for resp in [(500, {"detail": "boom"}), httpx.ReadTimeout("slow"), (200, {"oops": 1})]:
            with FakeCrm(resp).patch():
                self.assertEqual(self.send(ctx, settings), (True, None))
        self.assertTrue(self.db.is_submitted(1, OLD))
        self.assertIsNone(self.db.get_crm_ticket(1, OLD, "residential"))

    def test_reply_404_creates_new_and_updates(self):
        self.db.save_crm_ticket(1, OLD, "residential", 5, 5, "bot")
        ctx, settings = self.ctx()
        crm = FakeCrm((404, {"detail": "Тикет #5 не найден"}), (200, {"ticket_number": 6, "ticket_id": 9, "created": True}))
        with crm.patch(), self.assertLogs("crm", "WARNING") as logs:
            self.send(ctx, settings)
        self.assertEqual([crm.field(0, "reply_to_ticket"), crm.field(1, "reply_to_ticket")], ["5", None])
        self.assertEqual(self.db.get_crm_ticket(1, OLD, "residential").ticket_number, 6)
        self.assertTrue(any("#5 не найден" in m for m in logs.output))


# ------------------------------------------------------------------ модуль CRM
class TestCrmModule(unittest.TestCase):
    def req(self, **kw):
        base = dict(text="t", chat_id=CHAT, source_message_id=1, author="a", meta={}, file_name="f.xlsx",
                    data=b"x")
        return TicketRequest(**{**base, **kw})

    def test_error_codes_mapped(self):
        cases = {401: "не принял", 404: "выключен", 409: "чат отключён", 413: "больше 20 МБ",
                 422: "отклонила данные", 503: "не настроен", 500: "кодом 500"}
        for code, words in cases.items():
            with FakeCrm((code, {"detail": f"d{code}"})).patch():
                res = asyncio.run(crm_ingest.push_ticket(CRM, self.req()))
            self.assertFalse(res.ok)
            self.assertEqual(res.status_code, code)
            self.assertIn(words.replace("не принял", "не приняла"), res.summary())
            self.assertIn(f"HTTP {code}: d{code}", res.error)

    def test_timeout_message(self):
        with FakeCrm(httpx.ReadTimeout("slow")).patch():
            res = asyncio.run(crm_ingest.push_ticket(CRM, self.req()))
        self.assertEqual(res.error, "CRM не ответила за 20 с")
        self.assertEqual(crm_ingest.TIMEOUT_SECONDS, 20.0)

    def test_big_file_skipped(self):
        with patch("crm_ingest.MAX_FILE_BYTES", 2):
            r = crm_ingest.build_request(CRM, self.req(data=b"xyz"))
        self.assertIsNone(r["files"])
        self.assertEqual(r["data"]["text"], "t\nфайл не приложен: слишком большой")
        self.assertEqual(crm_ingest.build_request(CRM, self.req())["data"]["text"], "t")
        self.assertEqual(len(crm_ingest.build_request(CRM, self.req())["files"]), 1)

    def test_size_boundary(self):
        """Порог с запасом: прикладываем только при len(data) < 20 000 000 (лимит CRM — 20 МиБ)."""
        self.assertEqual(crm_ingest.MAX_FILE_BYTES, 20_000_000)
        for size, attached in ((19_999_999, True), (20_000_000, False), (20 * 1024 * 1024, False)):
            r = crm_ingest.build_request(CRM, self.req(data=b"x" * size))
            self.assertEqual(r["files"] is not None, attached, size)
            self.assertEqual(r["data"]["text"].endswith("файл не приложен: слишком большой"), not attached, size)

    def test_404_kinds(self):
        for detail, missing, words in (("Тикет #7 не найден", True, "тикет для дописывания не найден"),
                                       ("Not Found", False, "выключен"), ("", False, "выключен")):
            self.assertEqual(crm_ingest.is_ticket_missing(detail), missing)
            with FakeCrm((404, {"detail": detail} if detail else {})).patch():
                res = asyncio.run(crm_ingest.push_ticket(CRM, self.req(reply_to_ticket=7)))
            self.assertIn(words, res.error)
        # «выключено» при reply_to_ticket — без повтора; «не найден» — повтор без reply_to_ticket
        crm = FakeCrm((404, {"detail": "Not Found"}))
        with crm.patch():
            res = asyncio.run(crm_ingest.push_with_fallback(CRM, self.req(reply_to_ticket=7)))
        self.assertEqual((len(crm.requests), res.ok, res.fallback), (1, False, False))
        crm = FakeCrm((404, {"detail": "Тикет #7 не найден"}), (200, {"ticket_number": 8, "ticket_id": 8,
                                                                       "created": True}))
        with crm.patch():
            res = asyncio.run(crm_ingest.push_with_fallback(CRM, self.req(reply_to_ticket=7)))
        self.assertEqual((len(crm.requests), res.ok, res.fallback, res.ticket_number), (2, True, True, 8))
        self.assertEqual([crm.field(0, "reply_to_ticket"), crm.field(1, "reply_to_ticket")], ["7", None])

    def test_config_from_env(self):
        with patch.dict(os.environ, {"CRM_INGEST_URL": "https://crm.example.test/x", "CRM_INGEST_TOKEN": ""}):
            cfg = crm_ingest.load_crm_config()
            self.assertFalse(cfg.enabled)
            self.assertIn("CRM_INGEST_TOKEN", crm_ingest.describe_config(cfg))
        with patch.dict(os.environ, {"CRM_INGEST_URL": "https://crm.example.test/x", "CRM_INGEST_TOKEN": "secret-x"}):
            cfg = crm_ingest.load_crm_config()
            self.assertTrue(cfg.enabled)
            self.assertNotIn("secret-x", crm_ingest.describe_config(cfg))

    def test_caption_without_tag(self):
        self.assertEqual(build_caption(TAG, 1, OLD), TAG + " " + build_caption("", 1, OLD))
        self.assertEqual(bot.build_caption, build_caption)


class TestTelegramSender(unittest.TestCase):
    def test_real_sender_builds_send_document(self):
        """Настоящий TelegramSender на python-telegram-bot, но HTTP-слой подменён (без сети)."""
        from telegram.request import HTTPXRequest
        seen = []

        async def fake_do_request(self_, url, method, request_data=None, **kw):
            seen.append((url.rsplit("/", 1)[-1], request_data))
            if url.endswith("getMe"):
                result = {"id": 1, "is_bot": True, "first_name": "b", "username": "fixture_bot"}
            else:
                result = {"message_id": 4242, "date": 0, "chat": {"id": CHAT, "type": "supergroup"}}
            return 200, json.dumps({"ok": True, "result": result}).encode()

        with patch.object(HTTPXRequest, "do_request", fake_do_request):
            mid = asyncio.run(admin.TelegramSender(TOKEN).send_document(CHAT, b"PK-data", "f.xlsx", "cap"))
        self.assertEqual(mid, 4242)
        method, rd = seen[-1]
        self.assertEqual(method, "sendDocument")
        self.assertEqual((rd.parameters["chat_id"], rd.parameters["caption"]), (CHAT, "cap"))
        self.assertEqual(rd.multipart_data["document"][:2], ("f.xlsx", b"PK-data"))


class TestMigration(Base):
    def test_new_tables_added_to_existing_db(self):
        """Старая БД без новых таблиц: init() их создаёт, остальные данные не трогает."""
        self.db.save_reading(1, "1", OLD, 100, None)
        with self.db.engine.begin() as conn:
            conn.execute(text("DROP TABLE chat_resends"))
            conn.execute(text("DROP TABLE crm_tickets"))
        url = str(self.db.engine.url)
        self.db.dispose()
        self.db = Database(url, base_dir=self.tmp)
        self.db.init()
        self.db.init()  # повторный запуск безопасен
        names = inspect(self.db.engine).get_table_names()
        self.assertIn("chat_resends", names)
        self.assertIn("crm_tickets", names)
        self.assertEqual(self.db.get_readings(1, OLD), {"1": 100})


if __name__ == "__main__":
    unittest.main()
