"""Офлайн-тесты (без Telegram): БД, генерация Excel, миграция архива, разбор ввода.

Запуск:  python -m unittest discover -s tests -v
"""
import asyncio
import io
import os
import shutil
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = ROOT / "tests" / "fixtures"
sys.path.insert(0, str(ROOT))
# Тесты работают на вымышленных домах и шаблонах из tests/fixtures (без настоящих адресов).
os.environ["HOUSES_FILE"] = str(FIXTURES / "houses.json")

from openpyxl import load_workbook  # noqa: E402

import config  # noqa: E402
from db import Database, Reading, Submission, previous_period  # noqa: E402
from excel_report import read_apartment_rows  # noqa: E402
from migrate_archive import migrate  # noqa: E402
from parsing import parse_readings  # noqa: E402


def make_db(tmp: Path) -> Database:
    for name in ("template_1.xlsx", "template_2.xlsx", "template_1_commerce.xlsx", "template_2_commerce.xlsx"):
        shutil.copy(FIXTURES / name, tmp / name)
    db = Database(f"sqlite:///{tmp / 'test.db'}", base_dir=tmp)
    db.init()
    return db


def wb_from_bytes(data: bytes):
    return load_workbook(io.BytesIO(data))


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.db = make_db(self.tmp)

    def tearDown(self):
        self.db.dispose()
        shutil.rmtree(self.tmp, ignore_errors=True)


class TestParsing(unittest.TestCase):
    def test_decimals_and_errors(self):
        r = parse_readings("1 100,5 2 200.25\n3 abc 4 400 5")
        self.assertEqual(r.entries, [("1", 100.5), ("2", 200.25), ("4", 400.0)])
        self.assertEqual(len(r.errors), 2)
        self.assertIn("3", r.errors[0])
        self.assertIn("5", r.errors[1])

    def test_letter_apartment(self):
        r = parse_readings("2Ц 150")
        self.assertEqual(r.entries, [("2ц", 150.0)])

    def test_negative_rejected(self):
        r = parse_readings("1 -5")
        self.assertEqual(r.entries, [])
        self.assertEqual(len(r.errors), 1)


class TestConfig(unittest.TestCase):
    def test_allowed_ids(self):
        self.assertEqual(config.parse_allowed_user_ids(" 1, 22 ,,333"), frozenset({1, 22, 333}))
        self.assertEqual(config.parse_allowed_user_ids(""), frozenset())
        with self.assertRaises(config.ConfigError):
            config.parse_allowed_user_ids("1,abc")

    def test_validation(self):
        env = dict(os.environ)
        try:
            os.environ["BOT_TOKEN"] = "bad"
            os.environ["GROUP_CHAT_ID"] = "-100123"
            with self.assertRaises(config.ConfigError):
                config.load_settings()
            os.environ["BOT_TOKEN"] = "123456:" + "A" * 35
            os.environ["GROUP_CHAT_ID"] = "x"
            with self.assertRaises(config.ConfigError):
                config.load_settings()
            os.environ["GROUP_CHAT_ID"] = "-100123"
            os.environ["DATABASE_URL"] = "sqlite:///:memory:"
            s = config.load_settings()
            self.assertEqual(s.group_chat_id, -100123)
        finally:
            os.environ.clear()
            os.environ.update(env)


class TestHousesFile(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _write(self, data):
        import json
        p = self.tmp / "houses.json"
        p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return p

    def test_fixture_loaded(self):
        self.assertEqual(config.HOUSES[1], {"name": "Дом 1", "hashtag": "#Дом1", "template": "template_1.xlsx",
                                            "commerce_template": "template_1_commerce.xlsx"})
        self.assertEqual(list(config.HOUSES), [1, 2])
        self.assertEqual(config.HOUSES_FILE, FIXTURES / "houses.json")

    def test_missing_file_message(self):
        with self.assertRaises(config.ConfigError) as cm:
            config.load_houses(self.tmp / "nope.json")
        self.assertIn("houses.example.json", str(cm.exception))

    def test_example_file_is_valid(self):
        houses, aliases = config.load_houses(ROOT / "houses.example.json")
        self.assertEqual([h["name"] for h in houses.values()], ["Дом 1", "Дом 2"])
        self.assertIn("*", aliases)

    def test_validation(self):
        good = {"name": "А", "hashtag": "#А", "template": "a.xlsx"}
        for bad in ({}, {"houses": {}}, {"houses": {"x": good}}, {"houses": {"1": {"name": "А"}}},
                    {"houses": {"1": good, "2": good}},                       # одинаковые названия
                    {"houses": {"1": good}, "aliases": {"7": {"а": "б"}}},  # псевдоним для несуществующего дома
                    {"houses": {"1": good}, "aliases": ["а"]}):
            with self.assertRaises(config.ConfigError, msg=bad):
                config.load_houses(self._write(bad))
        houses, aliases = config.load_houses(self._write({"_c": 1, "houses": {"2": good, "1": dict(good, name="Б")}}))
        self.assertEqual(list(houses), [1, 2])          # по номеру дома
        self.assertNotIn("commerce_template", houses[2])
        self.assertEqual(aliases, {})

    def test_env_override(self):
        env = dict(os.environ)
        try:
            os.environ["HOUSES_FILE"] = "/tmp/elsewhere.json"
            self.assertEqual(config.houses_file_path(), Path("/tmp/elsewhere.json"))
            os.environ["HOUSES_FILE"] = "conf/h.json"
            self.assertEqual(config.houses_file_path(), ROOT / "conf" / "h.json")
        finally:
            os.environ.clear()
            os.environ.update(env)


class TestDatabase(Base):
    def test_seed(self):
        a1 = self.db.apartment_map(1)
        a2 = self.db.apartment_map(2)
        self.assertEqual(len(a1), len(read_apartment_rows(FIXTURES / "template_1.xlsx")))
        self.assertIn("2ц", a1)
        self.assertNotIn("*сословабонента", a1)
        self.assertEqual(len(a2), 118)
        self.assertEqual(a1["1"].row_order, 5)
        self.db.sync_houses()  # повторно — без дублей
        self.assertEqual(len(self.db.apartment_map(1)), len(a1))

    def test_aliases_seeded_once(self):
        names = {a.alias for a in self.db.list_aliases()}
        self.assertEqual(names, set(config.DEFAULT_ALIASES["*"]))
        self.db.delete_alias(next(a.id for a in self.db.list_aliases() if a.alias == "лютик"))
        self.db.init()  # перезапуск: удалённый в админке псевдоним не возвращается, дублей нет
        self.assertEqual({a.alias for a in self.db.list_aliases()}, names - {"лютик"})
        self.assertEqual(self.db.seed_aliases({"*": {"новый": "ромашка"}}), 0)  # рабочую БД не трогаем

    def test_readings_previous_month_and_update(self):
        self.db.save_reading(1, "1", "2026-09", 1000, 11)
        r = self.db.save_reading(1, "1", "2026-10", 900, 11)
        self.assertTrue(r.decreased)
        self.assertEqual(r.previous, 1000)
        r = self.db.save_reading(1, "1", "2026-10", 1100.5, 12)
        self.assertTrue(r.updated)
        self.assertFalse(r.decreased)
        self.assertIsNone(self.db.save_reading(1, "999", "2026-10", 1, 11))
        with self.db.Session() as s:
            rows = s.query(Reading).filter_by(period="2026-10").all()
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].value, 1100.5)
            self.assertEqual(rows[0].entered_by, 12)
        self.assertEqual(previous_period("2026-01"), "2025-12")

    def test_session_survives_restart(self):
        self.db.start_session(42, 2, "2026-10")
        self.db.dispose()
        db2 = Database(f"sqlite:///{self.tmp / 'test.db'}", base_dir=self.tmp)
        db2.init()
        s = db2.get_session(42)
        self.assertEqual((s.house_id, s.period), (2, "2026-10"))
        db2.end_session(42)
        self.assertIsNone(db2.get_session(42))
        db2.dispose()


class TestExcel(Base):
    def test_generation(self):
        self.db.save_reading(1, "1", "2026-09", 13000, 1)
        self.db.save_reading(1, "1", "2026-10", 13100.5, 1)
        self.db.save_reading(1, "2", "2026-10", 30300, 1)
        self.db.save_reading(1, "2ц", "2026-09", 50, 1)
        ws = wb_from_bytes(self.db.build_excel(1, "2026-10")).active
        self.assertEqual(ws["D4"].value, "Сентябрь")
        self.assertEqual(ws["E4"].value, "Октябрь")
        self.assertEqual(ws["D2"].value, datetime(2026, 10, 21))
        self.assertEqual((ws["D5"].value, ws["E5"].value, ws["F5"].value), (13000, 13100.5, 100.5))
        self.assertEqual((ws["D6"].value, ws["E6"].value, ws["F6"].value), (None, 30300, None))
        self.assertEqual((ws["D7"].value, ws["E7"].value), (None, None))  # старые значения шаблона стёрты
        self.assertEqual(ws["D75"].value, 50)
        self.assertEqual(ws["A78"].value, "*Со слов абонента")  # строки после квартир не тронуты
        self.assertEqual(ws["B5"].value, "Иванов И. И.")

    def test_redo_same_month_uses_previous_month(self):
        """Баг старой версии: повторная сессия в том же месяце брала текущие значения как предыдущие."""
        self.db.save_reading(2, "1", "2026-09", 100, 1)
        self.db.save_reading(2, "1", "2026-10", 150, 1)
        self.db.mark_submitted(2, "2026-10", 1, "f1.xlsx")
        self.db.save_reading(2, "1", "2026-10", 160, 1)  # переделали сессию
        ws = wb_from_bytes(self.db.build_excel(2, "2026-10")).active
        self.assertEqual((ws["D5"].value, ws["E5"].value, ws["F5"].value), (100, 160, 60))

    def test_january(self):
        ws = wb_from_bytes(self.db.build_excel(2, "2027-01")).active
        self.assertEqual((ws["D4"].value, ws["E4"].value), ("Декабрь", "Январь"))


def make_archive_file(path: Path, house: int, d2: datetime, values: dict, prev_values=None):
    wb = load_workbook(io.BytesIO((FIXTURES / f"template_{house}.xlsx").read_bytes()))
    ws = wb.active
    ws["D2"] = d2
    for row in range(5, ws.max_row + 1):
        apt = ws.cell(row=row, column=1).value
        key = str(apt).strip().lower() if apt is not None else None
        ws.cell(row=row, column=5).value = values.get(key)
        if prev_values is not None:
            ws.cell(row=row, column=4).value = prev_values.get(key)
    wb.save(path)
    wb.close()


class TestMigration(Base):
    def test_import_idempotent(self):
        arch = self.tmp / "archive"
        arch.mkdir()
        make_archive_file(arch / "20260921_1200_house1_Дом_1.xlsx", 1,
                          datetime(2026, 9, 21), {"1": 12900, "2": 30100})
        # Повтор в том же месяце: более поздний файл должен победить
        make_archive_file(arch / "20260922_0900_house1_Дом_1.xlsx", 1,
                          datetime(2026, 9, 21), {"1": 13000, "3": 22300})
        make_archive_file(arch / "20261021_1000_house2_Дом_2.xlsx", 2,
                          datetime(2026, 10, 21), {"1": 14300.5})
        (arch / "random.xlsx").write_bytes(b"not an xlsx")
        self.db.save_reading(2, "1", "2026-10", 14400, 77)  # введено ботом — не трогать

        st1 = migrate(self.db, arch, use_templates=True)
        self.assertEqual(st1["files"], 3)
        self.assertEqual(st1["files_failed"], 0)
        self.assertEqual(st1["submissions_added"], 3)
        self.assertGreater(st1["template_inserted"], 0)

        sep = self.db.get_readings(1, "2026-09")
        self.assertEqual(sep, {"1": 13000, "2": 30100, "3": 22300})
        self.assertEqual(self.db.get_readings(2, "2026-10")["1"], 14400)
        # Значения из шаблонов: дом 1 — октябрь 2025 (D), дом 2 — сентябрь/октябрь 2025
        self.assertEqual(self.db.get_readings(1, "2025-10")["1"], 13060)
        self.assertEqual(self.db.get_readings(2, "2025-10")["1"], 14173)
        self.assertEqual(self.db.get_readings(2, "2025-09")["1"], 14078)

        with self.db.Session() as s:
            n_readings = s.query(Reading).count()
            n_subs = s.query(Submission).count()
        st2 = migrate(self.db, arch, use_templates=True)
        self.assertEqual((st2["inserted"], st2["updated"], st2["submissions_added"], st2["template_inserted"]),
                         (0, 0, 0, 0))
        with self.db.Session() as s:
            self.assertEqual(s.query(Reading).count(), n_readings)
            self.assertEqual(s.query(Submission).count(), n_subs)

        # После миграции октябрьская ведомость дома 1 берёт «предыдущие» из сентябрьского архива
        ws = wb_from_bytes(self.db.build_excel(1, "2026-10")).active
        self.assertEqual(ws["D5"].value, 13000)
        self.assertTrue(self.db.is_submitted(1, "2026-09"))


class FakeTelegram:
    """Поддельные объекты Telegram для обработчиков бота."""

    def _ctx(self, send_side_effect=None):
        bot = SimpleNamespace(send_document=AsyncMock(side_effect=send_side_effect))
        settings = config.Settings(bot_token="x", group_chat_id=-100, allowed_user_ids=frozenset())
        self.db.seed_roles([], [5])  # пользователь 5 — обычный пользователь бота
        app = SimpleNamespace(bot_data={"db": self.db, "settings": settings})
        return SimpleNamespace(application=app, bot=bot, user_data={})

    def _open(self, ctx, house, kind="Жилые", user_id=5):
        """Выбор дома и ведомости (Жилые/Коммерция). Возвращает update последнего шага."""
        import bot
        asyncio.run(bot.handle_message(self._update(house, user_id), ctx))
        u = self._update(kind, user_id)
        asyncio.run(bot.handle_message(u, ctx))
        return u

    def _update(self, text, user_id=5):
        msg = SimpleNamespace(text=text, reply_text=AsyncMock(), reply_document=AsyncMock())
        return SimpleNamespace(message=msg, effective_user=SimpleNamespace(id=user_id, username="u"))



class TestBotHandlers(FakeTelegram, Base):
    def test_flow(self):
        import bot
        ctx = self._ctx(send_side_effect=[RuntimeError("network"), None])
        run = asyncio.run
        self._open(ctx, "Дом 2")
        sess = self.db.get_session(5)
        self.assertEqual(sess.house_id, 2)

        u = self._update("1 100,5 2 abc 999 5")
        run(bot.handle_message(u, ctx))
        reply = u.message.reply_text.call_args[0][0]
        self.assertIn("✅ Сохранено 1 квартира", reply)
        self.assertIn("2", reply)
        self.assertIn("999", reply)
        self.assertEqual(self.db.get_readings(2, sess.period), {"1": 100.5})

        u = self._update("Файл сейчас")
        run(bot.handle_message(u, ctx))
        self.assertEqual(u.message.reply_document.call_args.kwargs["caption"], "Текущий файл")

        u = self._update("Готово")
        run(bot.handle_message(u, ctx))  # первая отправка падает
        self.assertIn("Ошибка отправки", u.message.reply_text.call_args[0][0])
        self.assertIsNotNone(self.db.get_session(5))
        self.assertFalse(self.db.is_submitted(2, sess.period))

        run(bot.handle_message(u, ctx))  # повтор успешен
        self.assertTrue(self.db.is_submitted(2, sess.period))
        self.assertIsNone(self.db.get_session(5))
        kwargs = ctx.bot.send_document.call_args.kwargs
        self.assertIn("#Дом2", kwargs["caption"])
        # в чат — понятное имя за отчётный месяц, в БД — служебное с отметкой времени
        self.assertEqual(kwargs["filename"], f"Дом_2_{bot.period_title(sess.period).replace(' ', '_')}.xlsx")
        with self.db.Session() as s:
            names = [x.file_name for x in s.query(Submission).filter_by(status="sent")]
        self.assertEqual(len(names), 1)
        self.assertRegex(names[0], r"^\d{8}_\d{6}_house2_Дом_2\.xlsx$")

    def test_access_denied(self):
        import bot
        ctx = self._ctx()
        u = self._update("Дом 2", user_id=6)  # роли нет
        asyncio.run(bot.handle_message(u, ctx))
        self.assertIn("нет доступа", u.message.reply_text.call_args[0][0])
        self.assertIn("6", u.message.reply_text.call_args[0][0])
        self.assertIsNone(self.db.get_session(6))


class TestApartmentNormalization(unittest.TestCase):
    def test_variants(self):
        from apartments import normalize_apartment as n
        for raw in ["2ц", "2Ц", "2 ц", "2-ц", "2.ц", "2/ц", "2c", "2C", "2u", "ц2", "Ц-2",
                    "цоколь 2", "Цоколь2", "2 цоколь", "2цоколь", "цок 2"]:
            self.assertEqual(n(raw), "2ц", raw)
        self.assertEqual(n("15a"), "15а")      # латинская a -> кириллическая
        self.assertEqual(n("12"), "12")
        self.assertEqual(n(12), "12")
        self.assertEqual(n(12.0), "12")
        self.assertEqual(n("кв. 5"), "5")
        self.assertIsNone(n("*Со слов абонента"))
        self.assertIsNone(n("Итого 5"))
        self.assertIsNone(n("abc"))

    def test_labels_and_suggestions(self):
        from apartments import (apartment_label, format_suggestions, is_nonstandard,
                                suggest_apartments)
        self.assertEqual(apartment_label("2ц"), "Цоколь 2")
        self.assertEqual(apartment_label("5"), "Кв. 5")
        self.assertTrue(is_nonstandard("2ц"))
        self.assertFalse(is_nonstandard("12"))
        existing = ["1", "2", "3", "4", "5", "2ц", "3ц", "4ц"]
        self.assertEqual(format_suggestions(suggest_apartments("5ц", existing)),
                         "Может, 4ц (Цоколь 4) или 5?")
        self.assertEqual(suggest_apartments("1ц", existing), ["2ц", "1"])
        self.assertEqual(suggest_apartments("99", existing), [])


class TestParsingMultiToken(unittest.TestCase):
    def test_glue(self):
        cases = {
            "2 ц 345": [("2ц", 345.0)],
            "цоколь 2 345": [("2ц", 345.0)],
            "2 цоколь 345": [("2ц", 345.0)],
            "2цоколь 345": [("2ц", 345.0)],
            "ц2 345": [("2ц", 345.0)],
            "1 100 ц 2 345 3 c 10,5 4 200.5": [("1", 100.0), ("2ц", 345.0), ("3ц", 10.5), ("4", 200.5)],
            "1 100 2 200": [("1", 100.0), ("2", 200.0)],
        }
        for text, expected in cases.items():
            r = parse_readings(text)
            self.assertEqual(r.entries, expected, text)
            self.assertEqual(r.errors, [], text)

    def test_missing_value(self):
        r = parse_readings("1 100 2 ц")
        self.assertEqual(r.entries, [("1", 100.0)])
        self.assertEqual(len(r.errors), 1)


class TestNonstandardDb(Base):
    def test_list_and_excel_numbers(self):
        self.assertEqual(self.db.nonstandard_apartments(1), ["2ц", "3ц", "4ц"])  # естественный порядок
        self.assertEqual(self.db.nonstandard_apartments(2), [])
        self.db.save_reading(1, "4ц", "2026-10", 4900, 1)
        ws = wb_from_bytes(self.db.build_excel(1, "2026-10")).active
        self.assertEqual((ws["A75"].value, ws["A76"].value, ws["A77"].value), ("2ц", "4ц", "3ц"))
        self.assertEqual(ws["E76"].value, 4900)

    def test_pending_survives_restart_and_schema_upgrade(self):
        self.db.start_session(7, 1, "2026-10")
        self.db.set_pending_apartment(7, "3ц")
        self.db.dispose()
        db2 = Database(f"sqlite:///{self.tmp / 'test.db'}", base_dir=self.tmp)
        db2.init()
        self.assertEqual(db2.get_session(7).pending_apartment, "3ц")
        db2.start_session(7, 1, "2026-10")  # новая сессия сбрасывает выбор
        self.assertIsNone(db2.get_session(7).pending_apartment)
        db2.dispose()

    def test_old_schema_gets_column(self):
        import sqlite3
        path = self.tmp / "old.db"
        con = sqlite3.connect(path)
        con.execute("CREATE TABLE user_sessions (user_id BIGINT PRIMARY KEY, house_id INTEGER, "
                    "period VARCHAR(7), started_at DATETIME, updated_at DATETIME)")
        con.execute("INSERT INTO user_sessions VALUES (1, 1, '2026-10', NULL, NULL)")
        con.commit()
        con.close()
        db = Database(f"sqlite:///{path}", base_dir=self.tmp)
        db.init()
        self.assertIsNone(db.get_session(1).pending_apartment)
        db.dispose()


class TestNonstandardBot(FakeTelegram, Base):
    def _last(self, u):
        return u.message.reply_text.call_args

    def test_hint_button_and_pending_flow(self):
        import bot
        run = asyncio.run
        ctx = self._ctx()
        u = self._open(ctx, "Дом 1")
        text = self._last(u)[0][0]
        self.assertIn("Цоколь 2, Цоколь 3, Цоколь 4", text)
        self.assertIn("«2ц 345» или «цоколь 2 345»", text)
        kb = self._last(u).kwargs["reply_markup"].keyboard
        self.assertIn(bot.BTN_NONSTD, [b.text for row in kb for b in row])
        self.assertEqual([b.text for b in kb[0]], [bot.BTN_FILE, bot.BTN_DONE])

        u = self._update(bot.BTN_NONSTD)
        run(bot.handle_message(u, ctx))
        labels = [b.text for row in self._last(u).kwargs["reply_markup"].keyboard for b in row]
        self.assertEqual(labels, ["Цоколь 2", "Цоколь 3", "Цоколь 4", bot.BTN_BACK])

        run(bot.handle_message(self._update("Цоколь 3"), ctx))
        self.assertEqual(self.db.get_session(5).pending_apartment, "3ц")

        u = self._update("abc")
        run(bot.handle_message(u, ctx))
        self.assertIn("введи число", self._last(u)[0][0])
        self.assertEqual(self.db.get_session(5).pending_apartment, "3ц")

        u = self._update("345,5")
        run(bot.handle_message(u, ctx))
        self.assertIn("Цоколь 3: сохранено 345,5", self._last(u)[0][0])
        period = self.db.get_session(5).period
        self.assertEqual(self.db.get_readings(1, period)["3ц"], 345.5)
        self.assertIsNone(self.db.get_session(5).pending_apartment)

        # Обычный ввод многословной формы и подсказка при опечатке
        u = self._update("цоколь 2 100 5ц 7 2 u 101")
        run(bot.handle_message(u, ctx))
        reply = self._last(u)[0][0]
        self.assertIn("✅ Сохранено 1 квартира, обновлено 1 квартира", reply)
        self.assertIn("Квартира 5ц не найдена. Может, 4ц (Цоколь 4) или 5?", reply)
        self.assertEqual(self.db.get_readings(1, period)["2ц"], 101)

    def test_back_and_house_without_nonstandard(self):
        import bot
        run = asyncio.run
        ctx = self._ctx()
        u = self._open(ctx, "Дом 2")
        self.assertNotIn("Нестандартные", self._last(u)[0][0])
        self.assertIs(self._last(u).kwargs["reply_markup"], bot.SESSION_KEYBOARD)
        u = self._update(bot.BTN_NONSTD)
        run(bot.handle_message(u, ctx))
        self.assertIn("нет нестандартных", self._last(u)[0][0])

        self._open(ctx, "Дом 1")
        run(bot.handle_message(self._update("Цоколь 4"), ctx))
        run(bot.handle_message(self._update(bot.BTN_BACK), ctx))
        self.assertIsNone(self.db.get_session(5).pending_apartment)
        # выбор, затем обычная строка из нескольких пар — выбор сбрасывается, пары сохраняются
        run(bot.handle_message(self._update("Цоколь 4"), ctx))
        run(bot.handle_message(self._update("1 13100 4ц 5000"), ctx))
        sess = self.db.get_session(5)
        self.assertIsNone(sess.pending_apartment)
        self.assertEqual(self.db.get_readings(1, sess.period), {"1": 13100, "4ц": 5000})


class TestPlural(unittest.TestCase):
    def test_forms(self):
        from apartments import apartments_count
        expected = {0: "квартир", 1: "квартира", 2: "квартиры", 4: "квартиры", 5: "квартир",
                    11: "квартир", 12: "квартир", 14: "квартир", 21: "квартира", 22: "квартиры",
                    25: "квартир", 101: "квартира", 111: "квартир", 112: "квартир", 123: "квартиры"}
        for n, form in expected.items():
            self.assertEqual(apartments_count(n), f"{n} {form}", n)

    def test_summary(self):
        import bot
        self.assertEqual(bot.saved_summary(3, 1), "✅ Сохранено 3 квартиры, обновлено 1 квартира")
        self.assertEqual(bot.saved_summary(0, 2), "✅ Обновлено 2 квартиры")
        self.assertEqual(bot.saved_summary(11, 0), "✅ Сохранено 11 квартир")
        self.assertIsNone(bot.saved_summary(0, 0))


def create_old_unique_schema(path: Path, rows: int = 2):
    """БД как у первой версии ветки: UNIQUE на submissions.file_name."""
    from sqlalchemy import (BigInteger, Column, DateTime, ForeignKey, Integer, MetaData, String,
                            Table, create_engine, insert)
    md = MetaData()
    Table("houses", md, Column("id", Integer, primary_key=True), Column("name", String(200), unique=True),
          Column("hashtag", String(100)), Column("template", String(255)))
    subs = Table("submissions", md, Column("id", Integer, primary_key=True, autoincrement=True),
                 Column("house_id", ForeignKey("houses.id"), index=True),
                 Column("period", String(7), index=True), Column("submitted_at", DateTime),
                 Column("submitted_by", BigInteger), Column("file_name", String(255), unique=True),
                 Column("source", String(20)))
    eng = create_engine(f"sqlite:///{path}")
    md.create_all(eng)
    with eng.begin() as c:
        c.execute(insert(md.tables["houses"]).values(id=1, name="Дом 1", hashtag="#В", template="t"))
        for i in range(rows):
            c.execute(insert(subs).values(house_id=1, period="2026-09", submitted_at=datetime(2026, 9, 21),
                                          submitted_by=None, file_name=f"old{i}.xlsx", source="archive"))
    eng.dispose()


class TestSubmissionUnique(Base):
    def _unique(self, db):
        return db._file_name_unique_constraints()

    def test_duplicate_names_allowed_in_new_db(self):
        self.assertEqual(self._unique(self.db), ([], []))
        self.db.mark_submitted(1, "2026-10", 1, "same.xlsx")
        self.db.mark_submitted(1, "2026-10", 1, "same.xlsx")
        with self.db.Session() as s:
            names = [x.file_name for x in s.query(Submission).all()]
        self.assertEqual(names, ["same.xlsx", "same.xlsx"])

    def test_migrates_old_sqlite_unique(self):
        path = self.tmp / "old_unique.db"
        create_old_unique_schema(path)
        db = Database(f"sqlite:///{path}", base_dir=self.tmp)
        self.assertNotEqual(self._unique(db), ([], []))   # до миграции UNIQUE есть
        db.init()
        self.assertEqual(self._unique(db), ([], []))       # после — нет
        with db.Session() as s:
            old = s.query(Submission).order_by(Submission.id).all()
            self.assertEqual([x.file_name for x in old], ["old0.xlsx", "old1.xlsx"])  # данные на месте
        db.mark_submitted(1, "2026-10", 5, "old0.xlsx")
        from sqlalchemy import inspect as sa_inspect
        idx = {i["name"] for i in sa_inspect(db.engine).get_indexes("submissions")}
        self.assertTrue({"ix_submissions_file_name", "ix_submissions_house_id"} <= idx)
        db.dispose()
        db2 = Database(f"sqlite:///{path}", base_dir=self.tmp)  # повторный запуск — ничего не ломает
        db2.init()
        with db2.Session() as s:
            self.assertEqual(s.query(Submission).count(), 3)
        db2.dispose()

    def test_suffix_fallback_if_unique_remains(self):
        path = self.tmp / "old_unique2.db"
        create_old_unique_schema(path, rows=0)
        db = Database(f"sqlite:///{path}", base_dir=self.tmp)
        db._rebuild_sqlite_submissions = lambda: None  # имитируем неудачную миграцию
        db.init()
        db.mark_submitted(1, "2026-10", 1, "a.xlsx")
        db.mark_submitted(1, "2026-10", 1, "a.xlsx")
        with db.Session() as s:
            self.assertEqual(sorted(x.file_name for x in s.query(Submission).all()), ["a.xlsx", "a.xlsx#2"])
        db.dispose()


class TestDoneAndPeriod(FakeTelegram, Base):
    def test_two_done_same_minute_both_recorded(self):
        import bot
        ctx = self._ctx()
        for _ in range(2):
            self._open(ctx, "Дом 2")
            asyncio.run(bot.handle_message(self._update("1 100"), ctx))
            u = self._update("Готово")
            asyncio.run(bot.handle_message(u, ctx))
            self.assertIn("✅ Файл улетел", u.message.reply_text.call_args[0][0])
        with self.db.Session() as s:
            self.assertEqual(s.query(Submission).count(), 2)

    def test_mark_submitted_failure_warns_user(self):
        import bot
        ctx = self._ctx()
        self._open(ctx, "Дом 2")

        def boom(*a, **k):
            raise RuntimeError("db is locked")
        self.db.mark_submitted = boom
        u = self._update("Готово")
        asyncio.run(bot.handle_message(u, ctx))
        reply = u.message.reply_text.call_args[0][0]
        self.assertIn("⚠️ Файл улетел в группу, но отметку об отправке", reply)
        self.assertIn("db is locked", reply)
        self.assertEqual(ctx.bot.send_document.await_count, 1)

    def test_restored_session_other_period_notice(self):
        import bot
        ctx = self._ctx()
        self.db.start_session(5, 2, "2026-09")  # как будто сессия осталась с прошлого месяца
        u = self._update("1 100")
        asyncio.run(bot.handle_message(u, ctx))
        texts = [c[0][0] for c in u.message.reply_text.call_args_list]
        self.assertEqual(texts[0], "ℹ️ Продолжаем сессию за Сентябрь 2026 — Дом 2")
        self.assertIn("Сохранено 1 квартира", texts[1])
        self.assertEqual(self.db.get_readings(2, "2026-09"), {"1": 100})  # период не изменился
        u = self._update("2 200")
        asyncio.run(bot.handle_message(u, ctx))  # второй раз уведомление не повторяется
        self.assertEqual(len(u.message.reply_text.call_args_list), 1)

    def test_no_notice_for_current_period(self):
        import bot
        ctx = self._ctx()
        self._open(ctx, "Дом 2")
        u = self._update("1 100")
        asyncio.run(bot.handle_message(u, ctx))
        self.assertEqual(len(u.message.reply_text.call_args_list), 1)


# ============================================================== Коммерция
HOUSE1_COMMERCE = ["Магазин «ПРОДУКТЫ»", "ООО «Ромашка»", "Альфа ЛТД"]
HOUSE2_COMMERCE = ["Альфа ЛТД 1 подъезд", "Альфа ЛТД 2 подъезд", "ООО РОМАШКА 1 подъезд",
                    "ООО РОМАШКА 2 подъезд", "АПТЕКА", "Добрая Пекарня"]


def key_of(db, house, title):
    return {t: k for k, t in db.subscribers(house)}[title]


class TestCommerceTemplates(Base):
    def test_parse_both_templates(self):
        from excel_report import read_rows, read_sheet_values
        r1 = read_rows(FIXTURES / "template_1_commerce.xlsx", "commerce")
        r2 = read_rows(FIXTURES / "template_2_commerce.xlsx", "commerce")
        self.assertEqual([t for _, _, t in r1], HOUSE1_COMMERCE)
        self.assertEqual([t for _, _, t in r2], HOUSE2_COMMERCE)
        self.assertEqual([r for _, r, _ in r1], [6, 7, 8])
        v1 = read_sheet_values(FIXTURES / "template_1_commerce.xlsx", "commerce")
        v2 = read_sheet_values(FIXTURES / "template_2_commerce.xlsx", "commerce")
        self.assertEqual((v1["prev_month"], v1["cur_month"]), (9, 10))
        self.assertEqual((v2["prev_month"], v2["cur_month"]), (8, 9))
        self.assertEqual(sorted(v2["cur"].values())[-1], 66469)

    def test_pick_sheet_with_data_not_active(self):
        from excel_report import read_rows
        wb = load_workbook(io.BytesIO((FIXTURES / "template_1_commerce.xlsx").read_bytes()))
        wb.active = 1  # пустой «Лист1» активен
        path = self.tmp / "active_empty.xlsx"
        wb.save(path)
        wb.close()
        self.assertEqual(len(read_rows(path, "commerce")), 3)

    def test_seed_subscribers_separately(self):
        self.assertEqual([t for _, t in self.db.subscribers(1)], HOUSE1_COMMERCE)
        self.assertEqual([t for _, t in self.db.subscribers(2)], HOUSE2_COMMERCE)
        self.assertEqual(len(self.db.apartment_numbers(1)), 73)   # жилые не изменились
        self.assertEqual(len(self.db.apartment_numbers(2)), 118)
        self.db.sync_houses()
        self.assertEqual(len(self.db.subscribers(2)), 6)

    def test_excel_keeps_formulas(self):
        pharmacy = key_of(self.db, 2, "АПТЕКА")
        self.db.save_reading(2, pharmacy, "2026-09", 60497, 1, kind="commerce")
        self.db.save_reading(2, pharmacy, "2026-10", 60600.5, 1, kind="commerce")
        ws = wb_from_bytes(self.db.build_excel(2, "2026-10", "commerce")).active
        self.assertEqual((ws["D5"].value, ws["E5"].value, ws["F5"].value), ("Сентябрь", "Октябрь", "Расход"
                                                                            if ws["F5"].value == "Расход"
                                                                            else "расход"))
        self.assertEqual(ws["D2"].value, datetime(2026, 10, 21))
        self.assertEqual((ws["A10"].value, ws["D10"].value, ws["E10"].value), ("АПТЕКА", 60497, 60600.5))
        self.assertEqual(ws["F10"].value, "=E10-D10")
        self.assertEqual(ws["F6"].value, "=E6-D6")
        self.assertEqual((ws["D6"].value, ws["E6"].value), (None, None))  # старые значения шаблона стёрты
        self.assertEqual(ws["A3"].value, "Коммерческие подключения")
        self.assertEqual(ws["A4"].value, "Наименование")
        # жилая ведомость не задета
        self.assertEqual(self.db.get_readings(2, "2026-10"), {})

    def test_template_baseline_migration(self):
        arch = self.tmp / "archive"
        arch.mkdir()
        migrate(self.db, arch, use_templates=True)
        self.assertEqual(self.db.get_readings(1, "2025-09", "commerce")[key_of(self.db, 1, "Альфа ЛТД")], 2069)
        self.assertEqual(self.db.get_readings(1, "2025-10", "commerce")[key_of(self.db, 1, "ООО «Ромашка»")], 673)
        self.assertEqual(self.db.get_readings(2, "2025-08", "commerce")[key_of(self.db, 2, "АПТЕКА")], 60330)
        self.assertEqual(self.db.get_readings(2, "2025-09", "commerce")[key_of(self.db, 2, "Добрая Пекарня")], 66469)
        st = migrate(self.db, arch, use_templates=True)
        self.assertEqual(st["template_inserted"], 0)

    def test_archive_name_with_seconds_and_commerce(self):
        from migrate_archive import collect_archive
        arch = self.tmp / "archive"
        arch.mkdir()
        shutil.copy(FIXTURES / "template_2_commerce.xlsx", arch / "20261021_101530_house2_commerce_Дом_2.xlsx")
        make_archive_file(arch / "20261021_101531_house2_Дом_2.xlsx", 2, datetime(2026, 10, 21), {"1": 5})
        files = collect_archive(arch)
        self.assertEqual([(f[1], f[3]) for f in files], [(2, "commerce"), (2, "residential")])
        self.assertEqual(files[0][2], datetime(2026, 10, 21, 10, 15, 30))

    def test_old_db_without_kind_columns(self):
        import sqlite3
        path = self.tmp / "pre_commerce.db"
        con = sqlite3.connect(path)
        con.executescript("""
            CREATE TABLE houses (id INTEGER PRIMARY KEY, name VARCHAR(200) UNIQUE, hashtag VARCHAR(100),
                                 template VARCHAR(255));
            CREATE TABLE apartments (id INTEGER PRIMARY KEY, house_id INTEGER REFERENCES houses(id),
                                     number VARCHAR(20), row_order INTEGER, active BOOLEAN,
                                     CONSTRAINT uq_apartment_house_number UNIQUE (house_id, number));
            CREATE TABLE readings (id INTEGER PRIMARY KEY, apartment_id INTEGER REFERENCES apartments(id),
                                   period VARCHAR(7), value FLOAT, entered_by BIGINT, source VARCHAR(20),
                                   created_at DATETIME, updated_at DATETIME,
                                   CONSTRAINT uq_reading_apartment_period UNIQUE (apartment_id, period));
            CREATE TABLE user_sessions (user_id BIGINT PRIMARY KEY, house_id INTEGER, period VARCHAR(7),
                                        pending_apartment VARCHAR(20), started_at DATETIME, updated_at DATETIME);
            CREATE TABLE submissions (id INTEGER PRIMARY KEY, house_id INTEGER, period VARCHAR(7),
                                      submitted_at DATETIME, submitted_by BIGINT, file_name VARCHAR(255),
                                      source VARCHAR(20));
            INSERT INTO houses VALUES (1, 'Дом 1', '#Дом1', 'template_1.xlsx');
            INSERT INTO apartments VALUES (1, 1, '1', 5, 1);
            INSERT INTO readings VALUES (1, 1, '2026-09', 13000, 1, 'bot', NULL, NULL);
            INSERT INTO user_sessions VALUES (9, 1, '2026-10', NULL, NULL, NULL);
        """)
        con.commit()
        con.close()
        db = Database(f"sqlite:///{path}", base_dir=self.tmp)
        db.init()
        self.assertEqual(db.get_readings(1, "2026-09"), {"1": 13000})
        self.assertEqual(db.get_session(9).kind, "residential")
        self.assertEqual([t for _, t in db.subscribers(1)], HOUSE1_COMMERCE)
        k = key_of(db, 1, "Альфа ЛТД")
        self.assertIsNotNone(db.save_reading(1, k, "2026-10", 5, 1, kind="commerce"))
        db.mark_submitted(1, "2026-10", 1, "x.xlsx", kind="commerce")
        self.assertTrue(db.is_submitted(1, "2026-10", "commerce"))
        self.assertFalse(db.is_submitted(1, "2026-10"))
        db.dispose()


class TestAliases(unittest.TestCase):
    def _m(self, titles, aliases=None):
        from ledgers import CommerceMatcher, subscriber_key
        aliases = config.DEFAULT_ALIASES["*"] if aliases is None else aliases
        return CommerceMatcher([(subscriber_key(t), t) for t in titles], aliases)

    def _titles(self, m, text):
        r = m.match_at(text.split(), 0)
        return [m.titles[k] for k in r.candidates] if r else None

    def test_aliases_and_typos(self):
        m1, m2 = self._m(HOUSE1_COMMERCE), self._m(HOUSE2_COMMERCE)
        for q in ["лютик", "Лютик", "лютикк", "лютик.", "ЛЮТИК", "ромашка", "ооо ромашка"]:
            self.assertEqual(self._titles(m1, q), ["ООО «Ромашка»"], q)
            self.assertEqual(self._titles(m2, q), ["ООО РОМАШКА 1 подъезд", "ООО РОМАШКА 2 подъезд"], q)
        self.assertEqual(self._titles(m1, "альфа"), ["Альфа ЛТД"])
        self.assertEqual(self._titles(m2, "альфаа"), ["Альфа ЛТД 1 подъезд", "Альфа ЛТД 2 подъезд"])
        self.assertEqual(self._titles(m1, "продукты"), ["Магазин «ПРОДУКТЫ»"])
        self.assertEqual(self._titles(m2, "аптека"), ["АПТЕКА"])
        self.assertEqual(self._titles(m2, "пекарня"), ["Добрая Пекарня"])
        self.assertEqual(self._titles(m2, "добрая"), ["Добрая Пекарня"])
        self.assertEqual(self._titles(m2, "добрая пекарня"), ["Добрая Пекарня"])
        for q in ["ц", "цоколь", "кв", "12", "2ц", "abc", "ооо"]:
            self.assertIsNone(self._titles(m2, q), q)

    def test_custom_alias_file(self):
        import json
        tmp = Path(tempfile.mkdtemp())
        try:
            data = json.loads((FIXTURES / "houses.json").read_text(encoding="utf-8"))
            data["aliases"] = {"_comment": "x", "*": {"еда": "продукты"}, "2": {"булочная": "пекарня"}}
            (tmp / "h.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            _, al = config.load_houses(tmp / "h.json")
            self.assertNotIn("_comment", al)
            self.assertEqual(self._titles(self._m(HOUSE1_COMMERCE, al["*"]), "еда"), ["Магазин «ПРОДУКТЫ»"])
            self.assertEqual(self._titles(self._m(HOUSE2_COMMERCE, al["2"]), "булочная"), ["Добрая Пекарня"])
            self.assertIn("лютик", config.DEFAULT_ALIASES["*"])  # стартовые — из houses.json
            (tmp / "bad.json").write_text("{oops", encoding="utf-8")
            with self.assertRaises(config.ConfigError):  # битый файл — понятная ошибка, а не тихий откат
                config.load_houses(tmp / "bad.json")
        finally:
            shutil.rmtree(tmp)

    def test_mixed_parse_and_entrances(self):
        m1, m2 = self._m(HOUSE1_COMMERCE), self._m(HOUSE2_COMMERCE)
        r = parse_readings("12 345 лютик 67 альфа 420 5ц 10", m1)
        self.assertEqual(r.entries, [("12", 345.0), ("5ц", 10.0)])
        self.assertEqual([(m1.titles[k], v) for k, v, _ in r.commerce], [("ООО «Ромашка»", 67), ("Альфа ЛТД", 420)])
        self.assertEqual(r.errors, [])
        r = parse_readings("лютик 1 67 лютик1 68 лютик 2п 69 альфа 2 подъезд 5 цоколь 2 345", m2)
        self.assertEqual([(m2.titles[k], v) for k, v, _ in r.commerce],
                         [("ООО РОМАШКА 1 подъезд", 67), ("ООО РОМАШКА 1 подъезд", 68),
                          ("ООО РОМАШКА 2 подъезд", 69), ("Альфа ЛТД 2 подъезд", 5)])
        self.assertEqual(r.entries, [("2ц", 345.0)])
        r = parse_readings("12 345 лютик 67 альфа 420", m2)
        self.assertEqual(r.entries, [("12", 345.0)])
        self.assertEqual([(a["raw"], a["value"], len(a["candidates"])) for a in r.ambiguous],
                         [("лютик", 67, 2), ("альфа", 420, 2)])
        self.assertEqual(parse_readings("1 100,5 2 200.25").entries, [("1", 100.5), ("2", 200.25)])


class TestCommerceBot(FakeTelegram, Base):
    def _restart(self, ctx):
        """Имитация перезапуска: новое подключение к той же БД и пустые bot_data/user_data."""
        self.db.dispose()
        self.db = Database(f"sqlite:///{self.tmp / 'test.db'}", base_dir=self.tmp)
        self.db.init()
        fresh = self._ctx()
        fresh.bot = ctx.bot
        return fresh

    def _labels(self, u):
        return [b.text for row in u.message.reply_text.call_args.kwargs["reply_markup"].keyboard for b in row]

    def test_commerce_session_flow(self):
        import bot
        run = asyncio.run
        ctx = self._ctx()
        u = self._update("Дом 2")
        run(bot.handle_message(u, ctx))
        self.assertEqual(self._labels(u), ["Жилые", "Коммерция", "Назад"])
        u = self._update("Коммерция")
        run(bot.handle_message(u, ctx))
        self.assertIn("Коммерция", u.message.reply_text.call_args[0][0])
        self.assertEqual(self._labels(u), HOUSE2_COMMERCE + ["Файл сейчас", "Готово", "Назад"])
        sess = self.db.get_session(5)
        self.assertEqual((sess.house_id, sess.kind), (2, "commerce"))
        pharmacy = key_of(self.db, 2, "АПТЕКА")
        from db import previous_period
        self.db.save_reading(2, pharmacy, previous_period(sess.period), 60497, 1, kind="commerce")

        u = self._update("АПТЕКА")
        run(bot.handle_message(u, ctx))
        self.assertIn("60497", u.message.reply_text.call_args[0][0])
        self.assertEqual(self.db.get_session(5).pending_apartment, pharmacy)

        ctx = self._restart(ctx)  # выбор абонента переживает перезапуск
        u = self._update("60500")
        run(bot.handle_message(u, ctx))
        reply = u.message.reply_text.call_args[0][0]
        self.assertIn("АПТЕКА: сохранено 60500 (расход 3)", reply)
        self.assertIn("✅ АПТЕКА", self._labels(u))
        self.assertIn("Добрая Пекарня", self._labels(u))

        u = self._update("АПТЕКА")       # кнопка без ✅ тоже работает, "✅ АПТЕКА" — тоже
        run(bot.handle_message(self._update("✅ АПТЕКА"), ctx))
        u = self._update("60400")
        run(bot.handle_message(u, ctx))
        self.assertIn("АПТЕКА: обновлено 60400", u.message.reply_text.call_args[0][0])
        self.assertIn("меньше предыдущего", u.message.reply_text.call_args[0][0])

        u = self._update("добрая 70000,5\nальфа 2 1700")
        run(bot.handle_message(u, ctx))
        self.assertIn("✅ Сохранено 2 абонента", u.message.reply_text.call_args[0][0])
        u = self._update("12 345")
        run(bot.handle_message(u, ctx))
        self.assertIn("Квартиры вводятся", u.message.reply_text.call_args[0][0])

        u = self._update("Готово")
        run(bot.handle_message(u, ctx))
        kw = ctx.bot.send_document.call_args.kwargs
        self.assertEqual(ctx.bot.send_document.await_count, 1)
        self.assertIn("— Коммерция", kw["caption"])
        self.assertIn("#Дом2 #Коммерция", kw["caption"])
        self.assertEqual(kw["filename"], f"Дом_2_Коммерция_{bot.period_title(sess.period).replace(' ', '_')}.xlsx")
        self.assertTrue(self.db.is_submitted(2, sess.period, "commerce"))
        self.assertFalse(self.db.is_submitted(2, sess.period, "residential"))
        self.assertIsNone(self.db.get_session(5))

    def test_residential_mixed_with_commerce_and_done_sends_both(self):
        import bot
        run = asyncio.run
        ctx = self._ctx()
        u = self._open(ctx, "Дом 1")
        self.assertIn("Коммерцию можно писать прямо здесь", u.message.reply_text.call_args[0][0])
        self.assertIn("укажи номер: «лютик 1 67»", u.message.reply_text.call_args[0][0])  # первый псевдоним
        u = self._update("12 345 лютик 67 альфа 420 4ц 10")
        run(bot.handle_message(u, ctx))
        reply = u.message.reply_text.call_args[0][0]
        self.assertIn("✅ Сохранено 2 квартиры", reply)
        self.assertIn("Коммерция: ООО «Ромашка» — 67", reply)
        self.assertIn("Коммерция: Альфа ЛТД — 420", reply)
        period = self.db.get_session(5).period
        self.assertEqual(self.db.get_readings(1, period), {"12": 345, "4ц": 10})
        com = self.db.get_readings(1, period, "commerce")
        self.assertEqual(com[key_of(self.db, 1, "ООО «Ромашка»")], 67)

        u = self._update("Файл сейчас")
        run(bot.handle_message(u, ctx))
        self.assertEqual(u.message.reply_document.await_count, 2)

        u = self._update("Готово")
        run(bot.handle_message(u, ctx))
        calls = ctx.bot.send_document.call_args_list
        self.assertEqual(len(calls), 2)
        self.assertNotIn("Коммерция", calls[0].kwargs["caption"])
        self.assertIn("#Коммерция", calls[1].kwargs["caption"])
        self.assertTrue(calls[0].kwargs["filename"].startswith("Дом_1_"))
        self.assertNotIn("Коммерция", calls[0].kwargs["filename"])
        self.assertTrue(calls[1].kwargs["filename"].startswith("Дом_1_Коммерция_"))
        self.assertIn("Коммерция тоже отправлена", u.message.reply_text.call_args[0][0])
        self.assertTrue(self.db.is_submitted(1, period, "commerce"))
        # Коммерческий файл содержит введённое, жилой — только квартиры
        ws = wb_from_bytes(self.db.build_excel(1, period, "commerce")).active
        self.assertEqual((ws["E7"].value, ws["E8"].value), (67, 420))

        # Следующая жилая сессия без новой коммерции — уходит только жилой файл
        self._open(ctx, "Дом 1")
        run(bot.handle_message(self._update("1 13100"), ctx))
        run(bot.handle_message(self._update("Готово"), ctx))
        self.assertEqual(ctx.bot.send_document.await_count, 3)

    def test_ambiguous_entrance_buttons_restart_safe(self):
        import bot
        run = asyncio.run
        ctx = self._ctx()
        self._open(ctx, "Дом 2")
        u = self._update("1 100 лютик 67 альфа 420")
        run(bot.handle_message(u, ctx))
        texts = [c[0][0] for c in u.message.reply_text.call_args_list]
        self.assertIn("✅ Сохранено 1 квартира", texts[0])
        self.assertIn("«лютик 67» — какой именно?", texts[1])
        self.assertEqual(self._labels(u), ["ООО РОМАШКА 1 подъезд", "ООО РОМАШКА 2 подъезд", "Назад"])
        period = self.db.get_session(5).period
        self.assertEqual(self.db.get_readings(2, period, "commerce"), {})  # ничего не сохранено

        u = self._update("Готово")
        run(bot.handle_message(u, ctx))
        self.assertIn("Сначала уточни", u.message.reply_text.call_args[0][0])
        self.assertEqual(ctx.bot.send_document.await_count, 0)

        ctx = self._restart(ctx)
        u = self._update("ООО РОМАШКА 2 подъезд")
        run(bot.handle_message(u, ctx))
        texts = [c[0][0] for c in u.message.reply_text.call_args_list]
        self.assertIn("Коммерция: ООО РОМАШКА 2 подъезд — 67", texts[0])
        self.assertIn("«альфа 420» — какой именно?", texts[1])
        u = self._update("1")  # номер подъезда текстом
        run(bot.handle_message(u, ctx))
        self.assertIn("Коммерция: Альфа ЛТД 1 подъезд — 420", u.message.reply_text.call_args[0][0])
        com = self.db.get_readings(2, period, "commerce")
        self.assertEqual(com, {key_of(self.db, 2, "ООО РОМАШКА 2 подъезд"): 67,
                               key_of(self.db, 2, "Альфа ЛТД 1 подъезд"): 420})
        self.assertEqual(self.db.get_pending_choices(5), [])

        # «Назад» пропускает неоднозначную запись
        run(bot.handle_message(self._update("лютик 70"), ctx))
        u = self._update("Назад")
        run(bot.handle_message(u, ctx))
        self.assertIn("Пропустил «лютик 70»", u.message.reply_text.call_args_list[0][0][0])
        self.assertEqual(self.db.get_pending_choices(5), [])

    def test_kind_choice_after_restart_and_back(self):
        import bot
        run = asyncio.run
        ctx = self._ctx()
        run(bot.handle_message(self._update("Дом 2"), ctx))
        ctx = self._restart(ctx)  # выбор дома (до «Жилые/Коммерция») хранится в памяти
        u = self._update("Коммерция")
        run(bot.handle_message(u, ctx))
        self.assertIn("Сначала выбери дом", u.message.reply_text.call_args[0][0])
        run(bot.handle_message(self._update("Дом 2"), ctx))
        u = self._update("Назад")
        run(bot.handle_message(u, ctx))
        self.assertEqual(u.message.reply_text.call_args[0][0], "Выбери дом:")

    def test_house_without_commerce_template_goes_straight_to_residential(self):
        import bot
        (self.tmp / "template_2_commerce.xlsx").unlink()
        ctx = self._ctx()
        u = self._update("Дом 2")
        asyncio.run(bot.handle_message(u, ctx))
        self.assertIn("Сессия начата", u.message.reply_text.call_args[0][0])
        self.assertEqual(self.db.get_session(5).kind, "residential")


if __name__ == "__main__":
    unittest.main()
