"""Тесты веб-админки (FastAPI TestClient) и связанной логики бота.

Запуск:  python -m unittest discover -s tests -v
"""
import asyncio
import io
import json
import re
import shutil
import sys
import time
import unittest
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
# Вымышленные дома из tests/fixtures — до импорта config (его читают admin и bot).
import os  # noqa: E402
os.environ["HOUSES_FILE"] = str(ROOT / "tests" / "fixtures" / "houses.json")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import admin  # noqa: E402
import bot  # noqa: E402
from admin_auth import AuthError, sign_login_widget, sign_webapp_init_data, validate_webapp_init_data  # noqa: E402
from config import AdminSettings  # noqa: E402
from db import LoginToken, ReadingAudit, current_period, previous_period, utcnow  # noqa: E402
from excel_report import pick_sheet, read_rows  # noqa: E402
from ledgers import norm_words  # noqa: E402
sys.path.insert(0, str(ROOT / "tests"))
from test_offline import Base, FakeTelegram, key_of  # noqa: E402

TOKEN = "123456:" + "B" * 35
ADMIN, USER, STRANGER = 100, 200, 300


def csrf_of(html: str) -> str:
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


class AdminBase(Base):
    def setUp(self):
        super().setUp()
        self.db.seed_roles([ADMIN], [USER])
        self.settings = AdminSettings(secret_key="s" * 40, bot_token=TOKEN)
        self.client = TestClient(admin.create_app(self.db, self.settings))

    def login(self, uid=ADMIN):
        raw = self.db.create_login_token(uid)
        r = self.client.get(f"/auth/tg?token={raw}", follow_redirects=False)
        self.assertEqual(r.status_code, 303)
        return r

    def csrf(self):
        return csrf_of(self.client.get("/").text)

    def post(self, url, data):
        return self.client.post(url, data={"csrf": self.csrf(), **data}, follow_redirects=False)


# ======================================================================== вход
class TestAuth(AdminBase):
    def test_login_required(self):
        for url in ("/", "/readings", "/users", "/aliases", "/submissions", "/missing", "/charts", "/download"):
            r = self.client.get(url, follow_redirects=False)
            self.assertEqual(r.status_code, 303, url)
            self.assertEqual(r.headers["location"], "/login")
        self.assertEqual(self.client.get("/api/charts/totals").status_code, 401)
        self.assertIn("Админ-панель", self.client.get("/login").text)

    def test_one_time_link_single_use(self):
        raw = self.db.create_login_token(ADMIN)
        r = self.client.get(f"/auth/tg?token={raw}", follow_redirects=False)
        self.assertEqual(r.headers["location"], "/")
        self.assertEqual(self.client.get("/").status_code, 200)
        cookie = r.headers["set-cookie"].lower()
        self.assertIn("httponly", cookie)
        self.assertIn("samesite=lax", cookie)
        other = TestClient(admin.create_app(self.db, self.settings))
        r = other.get(f"/auth/tg?token={raw}", follow_redirects=False)  # второй раз — нельзя
        self.assertIn("/login?error=", r.headers["location"])
        self.assertEqual(other.get("/", follow_redirects=False).status_code, 303)
        with self.db.Session() as s:  # хранится только хеш
            self.assertIsNone(s.get(LoginToken, raw))

    def test_one_time_link_expired(self):
        raw = self.db.create_login_token(ADMIN, ttl_seconds=300)
        with patch("db.utcnow", return_value=utcnow() + timedelta(seconds=301)):
            self.assertIsNone(self.db.consume_login_token(raw))
        raw = self.db.create_login_token(ADMIN, ttl_seconds=300)
        with patch("db.utcnow", return_value=utcnow() + timedelta(seconds=299)):
            self.assertEqual(self.db.consume_login_token(raw), ADMIN)
        self.assertIsNone(self.db.consume_login_token("garbage"))

    def test_non_admin_denied(self):
        raw = self.db.create_login_token(USER)
        r = self.client.get(f"/auth/tg?token={raw}", follow_redirects=False)
        self.assertIn("/login?error=", r.headers["location"])
        self.assertEqual(self.client.get("/", follow_redirects=False).status_code, 303)
        init = sign_webapp_init_data({"auth_date": str(int(time.time())), "user": json.dumps({"id": USER})}, TOKEN)
        r = self.client.post("/auth/webapp", json={"init_data": init})
        self.assertEqual(r.status_code, 403)
        self.assertIn("Нет доступа", r.json()["error"])

    def test_role_revoked_kills_session(self):
        self.login()
        self.db.seed_roles([], [])
        self.db.set_role(USER, "admin", "test")
        self.db.set_role(ADMIN, "user", "test")  # больше не админ
        self.assertEqual(self.client.get("/", follow_redirects=False).status_code, 303)

    def test_webapp_init_data_with_signature(self):
        # Реальный Telegram присылает поле signature; оно участвует в HMAC наравне с остальными.
        now = int(time.time())
        fields = {"auth_date": str(now), "query_id": "AAH", "signature": "c2lnbmF0dXJl-_Ed25519",
                  "user": json.dumps({"id": ADMIN, "first_name": "A"})}
        good = sign_webapp_init_data(fields, TOKEN)
        self.assertEqual(validate_webapp_init_data(good, TOKEN)["id"], ADMIN)
        with self.assertRaises(AuthError):  # подмена signature ломает подпись
            validate_webapp_init_data(good.replace("c2lnbmF0dXJl", "c2lnbmF0dXJm"), TOKEN)
        r = self.client.post("/auth/webapp", json={"init_data": good})
        self.assertEqual(r.json(), {"ok": True, "redirect": "/"})

    def test_webapp_init_data(self):
        now = int(time.time())
        fields = {"auth_date": str(now), "query_id": "AAH", "user": json.dumps({"id": ADMIN, "first_name": "A"})}
        good = sign_webapp_init_data(fields, TOKEN)
        self.assertEqual(validate_webapp_init_data(good, TOKEN)["id"], ADMIN)
        with self.assertRaises(AuthError):
            validate_webapp_init_data(good, "999:" + "C" * 35)            # чужой токен
        with self.assertRaises(AuthError):
            validate_webapp_init_data(good.replace("AAH", "AAX"), TOKEN)  # подмена поля
        old = sign_webapp_init_data({**fields, "auth_date": str(now - 86400 - 10)}, TOKEN)
        with self.assertRaises(AuthError):
            validate_webapp_init_data(old, TOKEN)                          # старше суток
        with self.assertRaises(AuthError):
            validate_webapp_init_data("", TOKEN)
        # через HTTP
        self.assertEqual(self.client.post("/auth/webapp", json={"init_data": old}).status_code, 403)
        self.assertEqual(self.client.post("/auth/webapp", json={"init_data": good + "x"}).status_code, 403)
        r = self.client.post("/auth/webapp", json={"init_data": good})
        self.assertEqual(r.json(), {"ok": True, "redirect": "/"})
        self.assertEqual(self.client.get("/").status_code, 200)

    def test_login_widget(self):
        p = sign_login_widget({"id": str(ADMIN), "first_name": "A", "auth_date": str(int(time.time()))}, TOKEN)
        bad = dict(p, first_name="B")
        r = self.client.get("/auth/widget", params=bad, follow_redirects=False)
        self.assertIn("/login?error=", r.headers["location"])
        r = self.client.get("/auth/widget", params=p, follow_redirects=False)
        self.assertEqual(r.headers["location"], "/")
        self.assertEqual(self.client.get("/").status_code, 200)

    def test_csrf_and_logout(self):
        self.login()
        r = self.client.post("/users/add", data={"user_id": "5", "role": "user", "csrf": "bad"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.db.user_role(5), "none")
        r = self.client.post("/logout", data={"csrf": "bad"})
        self.assertEqual(r.status_code, 403)
        r = self.post("/logout", {})
        self.assertEqual(r.headers["location"], "/login")
        self.assertEqual(self.client.get("/", follow_redirects=False).status_code, 303)

    def test_secure_cookie_with_https(self):
        settings = AdminSettings(secret_key="s" * 40, bot_token=TOKEN, public_url="https://admin.example.org")
        c = TestClient(admin.create_app(self.db, settings), base_url="https://testserver")
        r = c.get(f"/auth/tg?token={self.db.create_login_token(ADMIN)}", follow_redirects=False)
        self.assertIn("secure", r.headers["set-cookie"].lower())


# ======================================================================== бот: кнопка и /admin
class TestBotAdminEntry(FakeTelegram, Base):
    def _labels(self, u):
        return [b.text for row in u.message.reply_text.call_args.kwargs["reply_markup"].keyboard for b in row]

    def test_button_only_for_admins(self):
        ctx = self._ctx()
        self.db.seed_roles([7], [5])
        u = self._update("/start", user_id=7)
        asyncio.run(bot.start(u, ctx))
        self.assertIn(bot.BTN_ADMIN, self._labels(u))
        u = self._update("/start", user_id=5)
        asyncio.run(bot.start(u, ctx))
        self.assertNotIn(bot.BTN_ADMIN, self._labels(u))

    def test_admin_command(self):
        ctx = self._ctx()
        ctx.application.bot_data["settings"].admin_public_url = "https://admin.example.org"
        self.db.seed_roles([7], [5])
        u = self._update("/admin", user_id=5)
        asyncio.run(bot.admin_command(u, ctx))
        self.assertEqual(u.message.reply_text.call_args[0][0], "Нет доступа")
        u = self._update(bot.BTN_ADMIN, user_id=5)   # кнопку набрали вручную
        asyncio.run(bot.handle_message(u, ctx))
        self.assertEqual(u.message.reply_text.call_args[0][0], "Нет доступа")

        u = self._update("/admin", user_id=7)
        asyncio.run(bot.admin_command(u, ctx))
        kb = u.message.reply_text.call_args.kwargs["reply_markup"].inline_keyboard
        self.assertEqual(kb[0][0].web_app.url, "https://admin.example.org/auth/webapp")
        link = kb[1][0].url
        token = link.split("token=")[1]
        self.assertTrue(link.startswith("https://admin.example.org/auth/tg?token="))
        self.assertEqual(self.db.consume_login_token(token), 7)
        self.assertIsNone(self.db.consume_login_token(token))

    def test_without_https_no_webapp_button(self):
        ctx = self._ctx()
        self.db.seed_roles([7], [])
        u = self._update("/admin", user_id=7)
        asyncio.run(bot.admin_command(u, ctx))
        kb = u.message.reply_text.call_args.kwargs["reply_markup"].inline_keyboard
        self.assertEqual(len(kb), 1)
        self.assertIsNone(kb[0][0].web_app)

    def test_roles_affect_bot_access_and_last_admin(self):
        ctx = self._ctx()
        self.db.seed_roles([7], [])
        u = self._update("Дом 2", user_id=8)
        asyncio.run(bot.handle_message(u, ctx))
        self.assertIn("нет доступа", u.message.reply_text.call_args[0][0])
        self.assertEqual(self.db.get_user(8).messages, 1)  # записан — виден в админке
        self.db.set_role(8, "user", "admin:7")
        self._open(ctx, "Дом 2", user_id=8)
        self.assertIsNotNone(self.db.get_session(8))
        self.db.set_role(8, "none", "admin:7")
        u = self._update("1 100", user_id=8)
        asyncio.run(bot.handle_message(u, ctx))
        self.assertIn("нет доступа", u.message.reply_text.call_args[0][0])
        from db import LastAdminError
        with self.assertRaises(LastAdminError):
            self.db.set_role(7, "user", "admin:7")
        self.db.set_role(8, "admin", "admin:7")
        self.db.set_role(7, "user", "admin:8")  # теперь можно
        self.assertTrue(self.db.is_admin(8))
        # роль, изменённая в админке, не перетирается .env при перезапуске
        self.db.seed_roles([7], [])
        self.assertEqual(self.db.user_role(7), "user")

    def test_deny_all_without_roles(self):
        self.db.seed_roles([], [])
        for uid in (1, 5, 42):
            u = self._update("Дом 2", user_id=uid)
            ctx = SimpleCtx(self.db)
            asyncio.run(bot.handle_message(u, ctx))
            self.assertIn("нет доступа", u.message.reply_text.call_args[0][0])
        self.assertEqual(self.db.count_active_users(), 0)


class SimpleCtx:
    def __init__(self, db):
        import config
        from types import SimpleNamespace
        self.application = SimpleNamespace(bot_data={"db": db, "settings": config.Settings(
            bot_token="x", group_chat_id=-100)})
        self.user_data = {}


def commerce(text, db, house):
    from parsing import parse_readings
    return parse_readings(text, matcher(db, house)).commerce


def matcher(db, house):
    return bot.get_matcher(SimpleCtx(db), house)


# ======================================================================== функции панели
class TestReadingsAdmin(AdminBase):
    def test_edit_with_audit_and_negative(self):
        self.login()
        p = current_period()
        apt = self.db.apartment_map(1)["1"]
        self.db.save_reading(1, "1", previous_period(p), 1000, USER)
        self.db.save_reading(1, "1", p, 1100, USER)
        r = self.post("/readings/save", {"apartment_id": apt.id, "period": p, "value": "900,5",
                                         "ret": f"/readings?house=1&period={p}"})
        self.assertEqual(r.headers["location"], f"/readings?house=1&period={p}")
        self.assertEqual(self.db.get_readings(1, p)["1"], 900.5)
        html = self.client.get(f"/readings?house=1&kind=residential&period={p}").text
        self.assertIn('class="neg', html)          # отрицательный расход подсвечен
        self.assertIn("-99,5", html)
        self.assertIn("меньше прошлого", html)     # flash
        self.assertIn(f"admin:{ADMIN}", html)
        log = self.db.audit_log([apt.id], p)
        self.assertEqual([(a.action, a.old_value, a.new_value, a.source) for a in log[:2]],
                         [("update", 1100, 900.5, "admin"), ("create", None, 1100, "bot")])
        self.assertEqual(log[0].actor, f"admin:{ADMIN}")
        self.assertEqual(log[1].actor, f"tg:{USER}")
        # создание и удаление
        apt2 = self.db.apartment_map(1)["2"]
        self.post("/readings/save", {"apartment_id": apt2.id, "period": p, "value": "50"})
        self.assertEqual(self.db.get_readings(1, p)["2"], 50)
        self.post("/readings/save", {"apartment_id": apt2.id, "period": p, "value": "abc"})
        self.assertEqual(self.db.get_readings(1, p)["2"], 50)
        self.post("/readings/delete", {"apartment_id": apt2.id, "period": p})
        self.assertNotIn("2", self.db.get_readings(1, p))
        with self.db.Session() as s:
            acts = [a.action for a in s.query(ReadingAudit).filter_by(apartment_id=apt2.id).order_by(ReadingAudit.id)]
        self.assertEqual(acts, ["create", "delete"])


class TestItemsAndRename(AdminBase):
    def test_add_disable_items(self):
        self.login()
        self.post("/items/add", {"house": 1, "kind": "residential", "number": "71", "title": "Новиков"})
        self.assertIn("71", self.db.apartment_map(1))
        apt = self.db.apartment_map(1)["71"]
        self.post(f"/items/{apt.id}/toggle", {})
        self.assertNotIn("71", self.db.apartment_map(1))
        self.post(f"/items/{apt.id}/toggle", {})
        self.assertIn("71", self.db.apartment_map(1))
        self.db.save_reading(1, "71", current_period(), 77, ADMIN)
        wb = load_workbook(io.BytesIO(self.db.build_excel(1, current_period())))
        ws = pick_sheet(wb)
        last = ws.max_row
        self.assertEqual((ws.cell(last, 1).value, ws.cell(last, 2).value, ws.cell(last, 5).value), ("71", "Новиков", 77))

    def test_rename_history_excel_and_reseed(self):
        self.login()
        key = key_of(self.db, 2, "АПТЕКА")
        apt = self.db.apartment_map(2, kind="commerce")[key]
        self.db.save_reading(2, key, "2026-05", 10, USER, kind="commerce")
        self.db.save_reading(2, key, "2026-06", 20, USER, kind="commerce")
        r = self.post(f"/items/{apt.id}/rename", {"new_name": "ООО Здоровье", "effective_period": "2026-06"})
        self.assertEqual(r.status_code, 303)
        hist = self.db.name_history(apt.id)
        self.assertEqual([(h.old_name, h.new_name, h.effective_period, h.changed_by) for h in hist],
                         [("АПТЕКА", "ООО Здоровье", "2026-06", f"admin:{ADMIN}")])
        page = self.client.get(f"/items/{apt.id}").text
        self.assertIn("ООО Здоровье", page)
        self.assertIn("История названий", page)

        def name_in_excel(period):
            ws = pick_sheet(load_workbook(io.BytesIO(self.db.build_excel(2, period, "commerce"))))
            return [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
        self.assertIn("АПТЕКА", name_in_excel("2026-05"))
        self.assertNotIn("ООО Здоровье", name_in_excel("2026-05"))
        self.assertIn("ООО Здоровье", name_in_excel("2026-06"))
        self.assertNotIn("АПТЕКА", name_in_excel("2026-06"))
        # бот распознаёт и старое, и новое имя
        m = matcher(self.db, 2)
        self.assertEqual(m.lookup(norm_words("здоровье")), [key])
        self.assertEqual(m.lookup(norm_words("аптека")), [key])

        # повторный посев: шаблон со старым именем и с новым — дублей нет
        count = len(self.db.list_items(2, "commerce"))
        self.db.sync_houses()
        self.assertEqual(len(self.db.list_items(2, "commerce")), count)
        tpl = self.tmp / "template_2_commerce.xlsx"
        row = next(r for k, r, t in read_rows(tpl, "commerce") if t == "АПТЕКА")
        wb = load_workbook(io.BytesIO(tpl.read_bytes()))
        pick_sheet(wb).cell(row, 1).value = "ООО Здоровье"
        wb.save(tpl)
        self.db.sync_houses()
        items = self.db.list_items(2, "commerce")
        self.assertEqual(len(items), count)
        self.assertEqual(sum(1 for i in items if i.active), count - sum(1 for i in items if not i.active))
        self.assertIn(key, self.db.apartment_map(2, kind="commerce"))
        self.assertEqual(self.db.get_readings(2, "2026-06", "commerce")[key], 20)


class TestAliasesAdmin(AdminBase):
    def test_crud_affects_bot_parser(self):
        from parsing import parse_readings
        self.login()
        pharmacy = key_of(self.db, 2, "АПТЕКА")
        self.assertEqual(commerce("маркет 100", self.db, 2), [])
        self.post("/aliases/save", {"house": "2", "alias": "маркет", "target": "аптека"})
        a = next(x for x in self.db.list_aliases() if x.alias == "маркет")
        self.assertEqual(commerce("маркет 100", self.db, 2), [(pharmacy, 100.0, "маркет")])
        self.assertEqual(commerce("маркет 100", self.db, 1), [])  # только дом 2
        html = self.client.get("/aliases").text
        self.assertIn("маркет", html)
        self.assertIn("Дом 2: АПТЕКА", html)
        r = self.post("/aliases/save", {"house": "2", "alias": "маркет", "target": "пекарня"})
        self.assertIn("уже", self.client.get("/aliases").text)  # дубль — ошибка
        self.post("/aliases/save", {"alias_id": a.id, "house": "2", "alias": "маркет", "target": "добрая"})
        bakery = key_of(self.db, 2, "Добрая Пекарня")
        self.assertEqual(commerce("маркет 5", self.db, 2), [(bakery, 5.0, "маркет")])
        self.post(f"/aliases/{a.id}/delete", {})
        self.assertEqual(commerce("маркет 5", self.db, 2), [])
        # стартовые псевдонимы из houses.json уже в БД
        self.assertIn("лютик", {x.alias for x in self.db.list_aliases()})


class TestUsersAdmin(AdminBase):
    def test_roles_page_pending_and_last_admin(self):
        self.login()
        self.db.touch_user(STRANGER, "stranger", "Иван", None)
        html = self.client.get("/users").text
        self.assertIn("Ожидают доступа", html)
        self.assertIn("@stranger", html)
        self.post(f"/users/{STRANGER}/role", {"role": "user"})
        self.assertEqual(self.db.user_role(STRANGER), "user")
        r = self.post(f"/users/{ADMIN}/role", {"role": "none"})
        self.assertEqual(self.db.user_role(ADMIN), "admin")
        self.assertIn("хотя бы один администратор", self.client.get("/users").text)
        self.post("/users/add", {"user_id": "555", "role": "admin"})
        self.assertTrue(self.db.is_admin(555))
        self.post(f"/users/{ADMIN}/role", {"role": "user"})   # теперь можно — второй админ есть
        self.assertEqual(self.db.user_role(ADMIN), "user")
        self.assertEqual(self.client.get("/users", follow_redirects=False).status_code, 303)  # сам потерял доступ


class TestSubmissionsMissingCharts(AdminBase):
    def test_submissions_and_download(self):
        self.login()
        p = current_period()
        self.db.save_reading(1, "1", p, 123, USER)
        self.db.mark_submitted(1, p, USER, "20261001_120000_house1_Дом_1.xlsx")
        self.db.mark_failed(1, p, USER, "x.xlsx", "network", kind="commerce")
        html = self.client.get("/submissions").text
        self.assertIn("20261001_120000_house1_Дом_1.xlsx", html)
        self.assertIn("ошибка", html)
        self.assertIn("network", html)
        r = self.client.get(f"/download?house=1&kind=residential&period={p}")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        title = admin.period_title(p).replace(" ", "_")
        from urllib.parse import quote
        self.assertIn(f"filename*=UTF-8''{quote(f'Дом_1_{title}.xlsx', safe='')}", r.headers["content-disposition"])
        ws = pick_sheet(load_workbook(io.BytesIO(r.content)))
        self.assertEqual(ws.cell(5, 5).value, 123)
        r2 = self.client.get(f"/download?house=1&kind=residential&period={p}")  # повторно — тоже можно
        self.assertEqual(r2.status_code, 200)

    def test_missing_report(self):
        self.login()
        p = current_period()
        total = len(self.db.apartment_map(2))
        self.db.save_reading(2, "1", p, 1, USER)
        self.db.save_reading(2, "2", p, 1, USER)
        d = self.client.get(f"/api/missing?period={p}").json()
        h2 = next(x for x in d["report"] if x["house"] == 2 and x["kind"] == "residential")
        self.assertEqual((h2["total"], h2["done"], len(h2["missing"])), (total, 2, total - 2))
        self.assertNotIn("Кв. 1", h2["missing"])
        self.assertIn("Кв. 3", h2["missing"])
        c = next(x for x in d["report"] if x["house"] == 2 and x["kind"] == "commerce")
        self.assertEqual(c["done"], 0)
        self.assertIn("АПТЕКА", c["missing"])
        html = self.client.get(f"/missing?period={p}").text
        self.assertIn("Кто не сдал", html)
        self.assertIn(f"не сдали <b class=\"danger\">{total - 2}</b>", html)

    def test_charts_endpoints(self):
        self.login()
        self.db.save_reading(1, "1", "2026-01", 100, USER)
        self.db.save_reading(1, "1", "2026-02", 150, USER)
        self.db.save_reading(1, "1", "2026-03", 140, USER)
        self.db.save_reading(1, "2", "2026-02", 10, USER)
        self.db.save_reading(1, "2", "2026-03", 30, USER)
        d = self.client.get("/api/charts/totals?house=1&kind=residential").json()
        self.assertEqual(d["periods"], ["2026-02", "2026-03"])
        self.assertEqual(d["totals"], [50, 10])
        self.assertEqual(d["counts"], [1, 2])
        self.assertEqual(d["labels"][0], "Февраль 2026")
        apt = self.db.apartment_map(1)["1"]
        d = self.client.get(f"/api/charts/item/{apt.id}").json()
        self.assertEqual(d["values"], [100, 150, 140])
        self.assertEqual(d["consumption"], [None, 50, -10])
        self.assertEqual(self.client.get("/api/charts/item/999999").status_code, 404)
        self.assertIn("/static/charts.js", self.client.get("/charts?house=1").text)
        self.assertEqual(self.client.get("/static/charts.js").status_code, 200)

    def test_pages_render(self):
        self.login()
        for url in ("/", "/readings?house=2&kind=commerce", "/items?house=2&kind=commerce", "/aliases", "/users",
                    "/submissions", "/missing", "/charts", "/login", "/auth/webapp"):
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
        self.assertEqual(self.client.get("/").headers["x-content-type-options"], "nosniff")


if __name__ == "__main__":
    unittest.main()


# ======================================================================== исправления после тестового развёртывания
class TestAliasAfterRename(FakeTelegram, AdminBase):
    """«ООО «Ромашка»» → «ИП Ромашкин»: псевдоним «лютик» (→ «ромашка») продолжает работать."""

    def _rename(self):
        key = key_of(self.db, 1, "ООО «Ромашка»")
        apt = self.db.apartment_map(1, kind="commerce")[key]
        self.db.rename_item(apt.id, "ИП Ромашкин", "2026-01", "admin:1")
        return key

    def _check_bot(self, key):
        ctx = self._ctx()
        self._open(ctx, "Дом 1")
        u = self._update("лютик 71")
        asyncio.run(bot.handle_message(u, ctx))
        reply = u.message.reply_text.call_args[0][0]
        self.assertNotIn("нет такой", reply)
        self.assertIn("ИП Ромашкин — 71", reply)
        self.assertEqual(self.db.get_readings(1, self.db.get_session(5).period, "commerce")[key], 71)

    def test_bound_alias_survives_rename(self):
        lex = next(a for a in self.db.list_aliases() if a.alias == "лютик")
        key = self._rename()
        self.assertIn(key, self.db.get_aliases(1)["лютик"][1])   # привязан по id
        self._check_bot(key)
        self.login()
        html = self.client.get("/aliases").text
        row = html[html.index("<b>лютик</b>"):]
        row = row[:row.index("</tr>")]
        self.assertIn("Дом 1: ИП Ромашкин", row)
        self.assertNotIn("никого", row)
        self.assertTrue(self.db.alias_targets()[lex.id])

    def test_text_alias_resolves_by_old_name_and_migrates(self):
        from db import AliasTarget
        with self.db.Session.begin() as s:        # как в старой БД: псевдонимы только текстом
            s.query(AliasTarget).delete()
        key = self._rename()
        self.assertEqual(self.db.get_aliases(1)["лютик"][1], [])
        self._check_bot(key)                        # фраза ищется и по прошлым названиям
        self.db.end_session(5)
        self.assertGreaterEqual(self.db.bind_aliases(), 1)  # автомиграция при init()
        self.assertIn(key, self.db.get_aliases(1)["лютик"][1])
        # повторная привязка ничего не дублирует
        self.assertEqual(self.db.bind_aliases(), 0)

    def test_new_alias_by_old_name_in_admin(self):
        self._rename()
        self.login()
        self.post("/aliases/save", {"house": "1", "alias": "ромка", "target": "ооо ромашка"})
        self.assertIn("Дом 1: ИП Ромашкин", self.client.get("/aliases").text)
        from parsing import parse_readings
        self.assertEqual([c[0] for c in parse_readings("ромка 5", matcher(self.db, 1)).commerce],
                         [key_of(self.db, 1, "ИП Ромашкин")])

    def test_matcher_old_name_unit(self):
        from ledgers import CommerceMatcher
        m = CommerceMatcher([("k1", "ИП Ромашкин"), ("k2", "Альфа ЛТД")], {"лютик": "ромашка"},
                            extra_names={"k1": ["ООО «Ромашка»"]})
        self.assertEqual(m.lookup(norm_words("лютик")), ["k1"])
        m = CommerceMatcher([("k1", "ИП Ромашкин")], {"лютик": ("что-то старое", ["k1", "чужой"])})
        self.assertEqual(m.lookup(norm_words("лютик")), ["k1"])


class TestHead(AdminBase):
    def test_head_on_get_routes(self):
        r = self.client.head("/login")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.content, b"")
        self.assertIn("text/html", r.headers["content-type"])
        self.assertEqual(self.client.head("/", follow_redirects=False).status_code, 303)
        self.assertEqual(self.client.head("/static/style.css").status_code, 200)
        self.assertEqual(self.client.head("/api/charts/totals").status_code, 401)
        self.login()
        for url in ("/", "/readings", "/users", "/aliases", "/submissions", "/missing", "/charts",
                    "/api/charts/totals", "/download?house=1&kind=residential"):
            r = self.client.head(url)
            self.assertEqual(r.status_code, 200, url)
            self.assertEqual(r.content, b"", url)
        self.assertEqual(self.client.head("/nope").status_code, 404)
        self.assertEqual(self.client.head("/logout").status_code, 405)  # только POST


class TestKeyboardsForAdmin(FakeTelegram, Base):
    def _labels(self, u):
        return [b.text for row in u.message.reply_text.call_args.kwargs["reply_markup"].keyboard for b in row]

    def test_choose_house_first_keeps_admin_button(self):
        ctx = self._ctx()
        self.db.seed_roles([7], [5])
        u = self._update("Жилые", user_id=7)            # ведомость без выбранного дома
        asyncio.run(bot.handle_message(u, ctx))
        self.assertEqual(u.message.reply_text.call_args[0][0], "Сначала выбери дом:")
        self.assertIn(bot.BTN_ADMIN, self._labels(u))
        u = self._update("Коммерция", user_id=5)
        asyncio.run(bot.handle_message(u, ctx))
        self.assertNotIn(bot.BTN_ADMIN, self._labels(u))

    def test_no_raw_house_keyboard_outside_helper(self):
        src = (ROOT / "bot.py").read_text(encoding="utf-8")
        body = src.split("def house_keyboard", 1)
        uses = [line for line in (body[0] + body[1].split("\n", 4)[-1]).splitlines()
                if "reply_markup=HOUSE_KEYBOARD" in line or "reply_markup=BASE_KEYBOARD" in line]
        self.assertEqual(uses, [])


class TestLoginTokenAtomic(Base):
    def test_concurrent_consumption_single_winner(self):
        import threading
        for _ in range(5):
            raw = self.db.create_login_token(ADMIN)
            results, barrier = [], threading.Barrier(8)

            def worker():
                barrier.wait()
                results.append(self.db.consume_login_token(raw))
            threads = [threading.Thread(target=worker) for _ in range(8)]
            for t in threads:
                t.start()
            for t in threads:
                t.join()
            self.assertEqual(results.count(ADMIN), 1, results)
            self.assertEqual(results.count(None), 7)

    def test_expired_and_used_flags(self):
        raw = self.db.create_login_token(ADMIN, ttl_seconds=-1)
        self.assertIsNone(self.db.consume_login_token(raw))
        with self.db.Session() as s:
            self.assertIsNone(s.query(LoginToken).one().used_at)   # просроченный не помечается
        raw = self.db.create_login_token(ADMIN)
        self.assertEqual(self.db.consume_login_token(raw), ADMIN)
        self.assertIsNone(self.db.consume_login_token(raw))


class TestCleanShutdown(unittest.TestCase):
    def _app(self, updater_stop_error=None):
        calls = []

        class Upd:
            running = False

            async def start_polling(self, **kw):
                calls.append("start_polling"); self.running = True

            async def stop(self):
                calls.append("updater.stop"); self.running = False
                if updater_stop_error:
                    raise updater_stop_error

        class App:
            running = False
            updater = Upd()

            async def initialize(self):
                calls.append("initialize")

            async def start(self):
                calls.append("start"); self.running = True

            async def stop(self):
                calls.append("stop"); self.running = False

            async def shutdown(self):
                calls.append("shutdown")

        return App(), calls

    def _run(self, app):
        import os
        import signal

        async def go():
            asyncio.get_running_loop().call_later(0.2, os.kill, os.getpid(), signal.SIGTERM)
            await bot.run_until_signal(app)
        asyncio.run(go())

    def test_sigterm_stops_cleanly(self):
        app, calls = self._app()
        self._run(app)
        self.assertEqual(calls, ["initialize", "start", "start_polling", "updater.stop", "stop", "shutdown"])

    def test_network_error_on_stop_still_shuts_down(self):
        from telegram.error import TimedOut
        app, calls = self._app(updater_stop_error=TimedOut())
        with self.assertLogs("__main__" if bot.__name__ == "__main__" else "bot", "WARNING") as logs:
            self._run(app)
        self.assertEqual(calls[-2:], ["stop", "shutdown"])
        self.assertIn("updater.stop", logs.output[0])


class TestMigrationLogMessage(Base):
    def test_unique_message_names_column(self):
        sys.path.insert(0, str(ROOT / "tests"))
        from test_offline import create_old_unique_schema
        from db import Database
        path = self.tmp / "old.db"
        create_old_unique_schema(path)
        db = Database(f"sqlite:///{path}", base_dir=self.tmp)
        with self.assertLogs("db", "WARNING") as logs:
            db.init()
        db.dispose()
        msg = next(m for m in logs.output if "UNIQUE" in m)
        self.assertIn("file_name", msg)
        self.assertNotIn("['']", msg)
        self.assertNotIn("()", msg)
