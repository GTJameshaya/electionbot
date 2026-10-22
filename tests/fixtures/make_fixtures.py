"""Генерирует вымышленные Excel-шаблоны для тестов — без настоящих имён, адресов и счетов.

Разметка повторяет настоящие шаблоны: жилые — месяцы в строке 4, данные с 5-й строки
(в доме 1 есть цоколь 2ц/4ц/3ц и сноска после квартир); коммерция — месяцы в строке 5,
данные с 6-й, в F формула =E-D, в доме 2 есть пары «1 подъезд / 2 подъезд».

Запуск (из корня проекта):  python tests/fixtures/make_fixtures.py
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

HERE = Path(__file__).resolve().parent
_THIN = Side(style="thin")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
TITLE = "Ведомость показаний электросчетчиков"


def _header(ws, house: str, d2: datetime) -> None:
    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = house
    ws["D2"] = d2
    ws["D2"].number_format = "DD.MM.YYYY"


def _boxed(ws, first: int, last: int) -> None:
    for r in range(first, last + 1):
        for c in range(1, 7):
            ws.cell(r, c).border = _BOX


def residential(name: str, house: str, d2: datetime, months: tuple[str, str], diff_title: str,
                rows: list[tuple], footnote: str | None = None, empty_sheet: bool = False) -> None:
    """rows: (номер квартиры, ФИО, прошлое, текущее). Формула расхода — только где есть прошлое."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ЭЛЕКТРИЧЕСТВО"
    _header(ws, house, d2)
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "№ кв", "ФИО", "Л/С", "Показания счетчиков"
    ws["D4"], ws["E4"], ws["F4"] = months[0], months[1], diff_title
    for rng in ("A1:F1", "A2:C2", "D2:F2", "A3:A4", "B3:B4", "C3:C4", "D3:F3"):
        ws.merge_cells(rng)
    r = 5
    for number, owner, prev, cur in rows:
        ws.cell(r, 1, number)
        ws.cell(r, 2, owner)
        ws.cell(r, 3, f"{900000 + r}/0000")
        if prev is not None:
            ws.cell(r, 4, prev)
            ws.cell(r, 6, f"=E{r}-D{r}")
        if cur is not None:
            ws.cell(r, 5, cur)
        r += 1
    _boxed(ws, 3, r - 1)
    if footnote:
        ws.cell(r, 1, footnote)
    if empty_sheet:
        wb.create_sheet("Лист1")
    wb.save(HERE / name)


def commerce(name: str, house: str, d2: datetime, months: tuple[str, str], rows: list[tuple],
             sheet_title: str = "ЭЛЕКТРИЧЕСТВО", empty_sheet: bool = False) -> None:
    """rows: (название абонента, прошлое, текущее)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    _header(ws, house, d2)
    ws["A3"] = "Коммерческие подключения"
    ws["A4"], ws["D4"] = "Наименование", "Показания счетчиков"
    ws["D5"], ws["E5"], ws["F5"] = months[0], months[1], "расход"
    for rng in ("A1:F1", "D2:F2", "A3:F3", "A4:C5", "D4:F4"):
        ws.merge_cells(rng)
    for i, (title, prev, cur) in enumerate(rows):
        r = 6 + i
        ws.cell(r, 1, title)
        ws.merge_cells(f"A{r}:C{r}")
        ws.cell(r, 4, prev)
        ws.cell(r, 5, cur)
        ws.cell(r, 6, f"=E{r}-D{r}")
    _boxed(ws, 3, 5 + len(rows))
    if empty_sheet:
        wb.create_sheet("Лист1")
    wb.save(HERE / name)


def main() -> None:
    # Дом 1: квартиры 1–70, потом цоколь в порядке 2ц, 4ц, 3ц и сноска. Заполнен только прошлый месяц.
    rows1 = [(n, "Иванов И. И." if n == 1 else f"Жилец {n}", 13060 if n == 1 else 10000 + n * 97, None)
             for n in range(1, 71)]
    rows1 += [("2ц", "Жилец 2ц", None, None), ("4ц", "Жилец 4ц", 4870, None), ("3ц", "Жилец 3ц", None, None)]
    residential("template_1.xlsx", "Дом 1", datetime(2025, 11, 21), ("Октябрь", "Ноябрь"), "расход",
                rows1, footnote="*Со слов абонента", empty_sheet=True)

    # Дом 2: квартиры 1–118, заполнены оба месяца.
    rows2 = [(n, f"Жилец {n}", 14078 if n == 1 else 5000 + n * 113, 14173 if n == 1 else 5050 + n * 113 + n % 40)
             for n in range(1, 119)]
    residential("template_2.xlsx", "Дом 2", datetime(2025, 10, 22), ("Сентябрь", "Октябрь"), "Расход", rows2)

    commerce("template_1_commerce.xlsx", "Дом 1", datetime(2025, 10, 21), ("Сентябрь", "Октябрь"),
             [("Магазин «ПРОДУКТЫ»", 25085, 25835), ("ООО «Ромашка»", 665, 673), ("Альфа ЛТД", 2069, 2088)],
             empty_sheet=True)
    commerce("template_2_commerce.xlsx", "Дом 2", datetime(2025, 9, 22), ("Август", "Сентябрь"),
             [("Альфа ЛТД 1 подъезд", 1437, 1455), ("Альфа ЛТД 2 подъезд", 1648, 1665),
              ("ООО РОМАШКА 1 подъезд", 3265, 3275), ("ООО РОМАШКА 2 подъезд", 10141, 10187),
              ("АПТЕКА", 60330, 60497), ("Добрая Пекарня", 64595, 66469)],
             sheet_title="Лист1")


if __name__ == "__main__":
    main()
