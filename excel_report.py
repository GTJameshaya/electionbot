"""Excel: читаем строки из шаблонов и собираем из БД готовую ведомость.

Разметка зависит от ведомости (ledgers.Ledger):
  * жилые — месяцы в строке 4 (D/E), данные с 5-й строки, в A номер квартиры, в F расход числом;
  * коммерция — месяцы в строке 5, данные с 6-й, в A название абонента, в F формула из шаблона.
В D2 — дата. Всё остальное в шаблоне остаётся как было.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from apartments import normalize_apartment
from ledgers import LEDGERS, RESIDENTIAL, Ledger, subscriber_key

FIRST_DATA_ROW = 5
COL_APT, COL_OWNER, COL_PREV, COL_CUR, COL_DIFF = 1, 2, 4, 5, 6  # A, B, D, E, F

MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
             "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

__all__ = ["normalize_apartment", "read_apartment_rows", "read_rows", "read_sheet_values", "build_workbook"]


def _load(path: Path, **kwargs):
    # Читаем файл сами — так дескриптор точно закроется.
    with open(path, "rb") as fh:
        return load_workbook(io.BytesIO(fh.read()), **kwargs)


def pick_sheet(wb):
    """Первый непустой лист. wb.active не годится — в шаблоне бывает пустой «Лист1»."""
    for ws in wb.worksheets:
        if any(c.value is not None for row in ws.iter_rows(min_row=1, max_row=12, max_col=6) for c in row):
            return ws
    return wb.worksheets[0]


def _ledger(ledger: Ledger | str | None) -> Ledger:
    if ledger is None:
        return LEDGERS[RESIDENTIAL]
    return LEDGERS[ledger] if isinstance(ledger, str) else ledger


def _rows(ws, ledger: Ledger) -> list[tuple[str, int, str]]:
    """[(ключ, строка, название), ...]. Название: у коммерции — из колонки A, у квартир — ФИО из B.
    Сноски (строки на «*») и повторы пропускаем."""
    rows, seen = [], set()
    for row in range(ledger.first_row, ws.max_row + 1):
        raw = ws.cell(row=row, column=COL_APT).value
        if ledger.key_mode == "apartment":
            owner = ws.cell(row=row, column=COL_OWNER).value
            key, title = normalize_apartment(raw), " ".join(str(owner).split()) if owner is not None else ""
        else:
            title = " ".join(str(raw).split()) if isinstance(raw, str) else ""
            key = subscriber_key(title) if title and not title.startswith("*") else None
        if key and key not in seen:
            seen.add(key)
            rows.append((key, row, title))
    return rows


def read_rows(path: Path, ledger: Ledger | str | None = None) -> list[tuple[str, int, str]]:
    """Строки ведомости из файла: [(ключ, строка, название), ...]."""
    wb = _load(path, data_only=True)
    try:
        return _rows(pick_sheet(wb), _ledger(ledger))
    finally:
        wb.close()


def read_apartment_rows(path: Path) -> list[tuple[str, int]]:
    """[(номер квартиры, номер строки), ...] из жилого шаблона или файла."""
    return [(k, r) for k, r, _ in read_rows(path, RESIDENTIAL)]


def _month_index(name) -> int | None:
    if not isinstance(name, str):
        return None
    name = name.strip().lower()
    for i, m in enumerate(MONTHS_RU):
        if m.lower() == name:
            return i + 1
    return None


def read_sheet_values(path: Path, ledger: Ledger | str | None = None) -> dict:
    """Достаёт данные из готовой ведомости (архивной или шаблона).

    Возвращает {"date": D2 или None, "prev_month": 1..12|None, "cur_month": 1..12|None,
                "prev": {ключ: значение}, "cur": {ключ: значение}}.
    """
    ledger = _ledger(ledger)
    wb = _load(path, data_only=True)
    try:
        ws = pick_sheet(wb)
        d2 = ws["D2"].value
        result = {
            "date": d2 if isinstance(d2, datetime) else None,
            "prev_month": _month_index(ws.cell(row=ledger.month_row, column=COL_PREV).value),
            "cur_month": _month_index(ws.cell(row=ledger.month_row, column=COL_CUR).value),
            "prev": {},
            "cur": {},
        }
        for key, row, _ in _rows(ws, ledger):
            for col, slot in ((COL_PREV, "prev"), (COL_CUR, "cur")):
                val = ws.cell(row=row, column=col).value
                if isinstance(val, bool):
                    continue
                if isinstance(val, (int, float)):
                    result[slot][key] = float(val)
                elif isinstance(val, str):
                    try:
                        result[slot][key] = float(val.strip().replace(",", "."))
                    except ValueError:
                        pass
        return result
    finally:
        wb.close()


def _num(value: float | None):
    """Число для ячейки: 12.0 → 12, дробные округляем до 3 знаков."""
    if value is None:
        return None
    return int(value) if float(value).is_integer() else round(float(value), 3)


@dataclass
class SheetRow:
    """Одна строка, которую надо записать в ведомость."""
    key: str
    row: int | None          # строка шаблона; None — добавлен в админке, строки в шаблоне нет
    prev: float | None
    cur: float | None
    name: str | None = None  # название/ФИО, действующее в периоде (None — не менять ячейку шаблона)
    label: str | None = None  # номер квартиры для добавленных строк (колонка A у жилых)


def _copy_style(ws, src_row: int, dst_row: int):
    from copy import copy
    for col in range(1, 7):
        src, dst = ws.cell(row=src_row, column=col), ws.cell(row=dst_row, column=col)
        if src.has_style:
            dst._style = copy(src._style)


def build_workbook(template_path: Path, year: int, month: int,
                   prev_values: dict[str, float] | None = None, cur_values: dict[str, float] | None = None,
                   ledger: Ledger | str | None = None, rows: list[SheetRow] | None = None) -> bytes:
    """Собирает xlsx (байты) по шаблону: D — прошлый месяц, E — текущий, F — расход.

    rows — строки из БД. Если их не передали, строки берутся из шаблона, а значения — из
    prev_values/cur_values. Старые D/E шаблона стираются; шапка, сноски и итоги не трогаются.
    У коммерции F не пишем — там формула =E-D. Строки, добавленные в админке, дописываются
    в конец листа со стилем первой строки данных.
    """
    ledger = _ledger(ledger)
    wb = _load(template_path)
    try:
        ws = pick_sheet(wb)
        prev_month = 12 if month == 1 else month - 1
        ws.cell(row=ledger.month_row, column=COL_PREV, value=MONTHS_RU[prev_month - 1])
        ws.cell(row=ledger.month_row, column=COL_CUR, value=MONTHS_RU[month - 1])
        ws["D2"] = datetime(year, month, 21)

        template_rows = _rows(ws, ledger)
        if rows is None:
            prev_values, cur_values = prev_values or {}, cur_values or {}
            rows = [SheetRow(k, r, prev_values.get(k), cur_values.get(k)) for k, r, _ in template_rows]
        for _, r, _ in template_rows:  # стираем старые значения шаблона
            ws.cell(row=r, column=COL_PREV).value = None
            ws.cell(row=r, column=COL_CUR).value = None
            if ledger.write_diff:
                ws.cell(row=r, column=COL_DIFF).value = None

        style_row = template_rows[0][1] if template_rows else ledger.first_row
        next_free = ws.max_row + 1
        for item in rows:
            r = item.row
            if not r:  # строки нет в шаблоне (добавлена в админке) — дописываем в конец
                r, next_free = next_free, next_free + 1
                _copy_style(ws, style_row, r)
                if ledger.key_mode == "apartment":
                    ws.cell(row=r, column=COL_APT).value = item.label or item.key
                if not ledger.write_diff:
                    ws.cell(row=r, column=COL_DIFF).value = f"=E{r}-D{r}"
            if item.name is not None:
                col = COL_APT if ledger.key_mode == "name" else COL_OWNER
                ws.cell(row=r, column=col).value = item.name
            ws.cell(row=r, column=COL_PREV).value = _num(item.prev)
            ws.cell(row=r, column=COL_CUR).value = _num(item.cur)
            if ledger.write_diff:
                diff = None
                if item.cur is not None and item.prev is not None:
                    diff = _num(round(item.cur - item.prev, 3))
                ws.cell(row=r, column=COL_DIFF).value = diff

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        wb.close()
